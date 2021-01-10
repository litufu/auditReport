# -*- coding: UTF-8 -*-
import pandas as pd
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from project.utils import checkLeftSpace, addTitle, addCombineTableTitle, addContentToCombineTitle, addParagraph, \
    createBorderedTable, addTable, setCell, addLandscapeContent, to_chinese, searchRecordItemByName, \
    handleName,searchModel,df_sort,addCombineTitleSpecialReceivable


# 第一步计算出报表项目编号
# 第二步显示所有编号的项目
# 根据条件过滤df
def filterDateFrame(sheetName, xlsxPath,conditions=("期末数","期初数")):
    df = pd.read_excel(xlsxPath, sheet_name=sheetName)
    if len(conditions)==0:
        return df
    s = False
    for condition in conditions:
        try:
            s = s | (df[condition].abs() > 0)
        except Exception as e:
            print(e)
            return df
    df1 = df[s]
    return df1

# 根据df生成报告
def dfToWord(document,df, style):
    dc = df.to_dict("split")
    addTable(document, dc, style=style)

def excelTableToWord(document, sheetName, xlsxPath, style,conditions=("期末数","期初数"),sort=False,sort_index_name="项目",sort_column_name="本期数",last_values=("其他","合计")):
    df = filterDateFrame(sheetName, xlsxPath, conditions)
    df = df.replace("—","  ", regex=True)
    if sort:
        df = df_sort(df,index_name=sort_index_name,column_name=sort_column_name,last_names=last_values)
    dc = df.to_dict("split")
    addTable(document, dc, style=style)



def addStart(document, context):
    companyType = context["report_params"]["companyType"]
    # 报告类型
    reportType = context["report_params"]["type"]
    # 报告日期
    reportDate = context["report_params"]["reportDate"]
    # 报告期间
    reportPeriod = context["report_params"]["reportPeriod"]
    lastPeriod = str.replace(reportPeriod, reportPeriod[:4], str(int(reportPeriod[:4]) - 1))
    # 人民币单位
    currencyUnit = context["notes_params"]["currencyUnit"]
    # 获取报告起始日
    startYear = reportDate[:4]

    if companyType == "国有企业":
        if reportType == "合并":
            addTitle(document, "八、合并财务报表重要项目的说明", 1, False)
        else:
            addTitle(document, "七、财务报表重要项目的说明", 1, False)
    else:
        if reportType == "合并":
            addTitle(document, "六、合并财务报表项目注释", 1, False)
        else:
            addTitle(document, "六、财务报表项目注释", 1, False)
    basicDesc = "以下注释项目除特别注明之外，金额单位为{}；期初数指{}年12月31日财务报表数，期末数指{}财务报表数，本期指{}，上期指{}".format(currencyUnit,
                                                                                          int(startYear) - 1,
                                                                                          reportDate, reportPeriod,
                                                                                          lastPeriod)
    addParagraph(document, basicDesc, "paragraph")

# （一）货币资金
def addMonetary(document, num, path, context):
    addTitle(document, "（{}）货币资金".format(to_chinese(num)), 2, True)
    df = filterDateFrame("货币资金", path)
    dfToWord(document,df,style=2)

    df = filterDateFrame("受限制的货币资金", path)
    if len(df)>0:
        addParagraph(document, "受限制的货币资金明细如下：", "paragraph")
        dfToWord(document, df, style=2)

# （二）交易性金融资产
def addTradingFinancialAssets(document, num, path, context):
    addTitle(document, "（{}）交易性金融资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "交易性金融资产", path, style=2)


# （三）以公允价值计量且其变动计入当期损益的金融资产
def addFinancialAssetsMeasuredAtFairValueWithChangesIncludedInCurrentProfitAndLoss(document, num, path, context):
    addTitle(document, "（{}）以公允价值计量且其变动计入当期损益的金融资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "以公允价值计量且其变动计入当期损益的金融资产", path, style=2)


# （四）衍生金融资产
def addDerivativeFinancialAssets(document, num, path, context):
    addTitle(document, "（{}）衍生金融资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "衍生金融资产", path, style=2)


# （五）应收票据
# 为表格添加合并标题，最后一列向上合并
def addCombineTitleSpecialLast(titles, table):
    for i, row in enumerate(titles):
        for j, cellText in enumerate(row):
            cell = table.cell(i, j)
            if cellText == "nan":
                if j >= 1 and j != len(row) - 1:
                    if checkLeftSpace(row, j):
                        cell.merge(table.cell(i - 1, j))
                    else:
                        cell.merge(table.cell(i, j - 1))
                else:
                    cell.merge(table.cell(i - 1, j))
            else:
                setCell(cell, cellText, WD_PARAGRAPH_ALIGNMENT.CENTER)


def addNotesReceivable(document, num, path, context):
    addTitle(document, "（{}）应收票据".format(to_chinese(num)), 2, True)
    # 老金融工具准则
    addParagraph(document, "1、应收票据分类", "paragraph")
    if context["notes_params"]["tandardssForFinancialInstruments"] == "老金融工具准则":
        excelTableToWord(document, "应收票据分类原金融工具准则", path, style=2)
    else:
        df = filterDateFrame("应收票据分类新金融工具准则",path,("期末账面余额","期初账面余额"))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["种类", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                      ["nan", "账面余额", "坏账准备", "账面价值", "账面余额", "坏账准备", "账面价值"]]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTableTitle(table, titles)
            addContentToCombineTitle(document, dc, table, titleLength, style=2)

    addParagraph(document, "2、期末已质押的应收票据", "paragraph")
    excelTableToWord(document, "已质押应收票据", path, 2,("期末已质押金额",))

    addParagraph(document, "3、期末已背书或贴现且在资产负债表日尚未到期的应收票据", "paragraph")
    excelTableToWord(document, "已背书或贴现且在资产负债表日尚未到期的应收票据", path, style=2,conditions=("期末终止确认金额","期末未终止确认金额"))

    addParagraph(document, "4、期末因出票人未履约而转为应收账款的票据", "paragraph")
    excelTableToWord(document, "因出票人未履约而转为应收账款的票据", path, style=2,conditions=("期末转应收账款金额",))

    if context["notes_params"]["tandardssForFinancialInstruments"] == "新金融工具准则":
        addParagraph(document, "5、期末单项计提坏账准备的应收票据", "paragraph")
        excelTableToWord(document, "期末单项计提坏账准备的应收票据新金融工具准则", path, style=2,conditions=("账面余额",))

        addParagraph(document, "6、采用组合计提坏账准备的应收票据", "paragraph")
        df = filterDateFrame("采用组合计提坏账准备的应收票据新金融工具准则",path,("期末账面余额",))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["项目", "期末数", "nan", "nan"],
                      ["nan", "账面余额", "坏账准备", "计提比例(%)"]]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTableTitle(table, titles)
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

        addParagraph(document, "7、坏账准备变动明细情况", "paragraph")
        df = filterDateFrame("应收票据坏账准备变动明细情况新金融工具准则",path)
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["项目", "期初数", "本期增加", "nan", "本期减少", "nan", "nan", "期末数"],
                      ["nan", "nan", "计提", "其他", "转回", "核销", "其他", "nan"]]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialLast(titles, table)
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

        addParagraph(document, "8、本期重要的坏账准备收回或转回情况", "paragraph")
        excelTableToWord(document, "本期重要的应收票据坏账准备收回或转回情况新金融工具准则", path, style=2,conditions=("收回或转回金额",))

        addParagraph(document, "9、本期实际核销的应收票据情况", "paragraph")
        excelTableToWord(document, "本期实际核销的应收票据情况新金融工具准则", path, style=2,conditions=("核销金额",))


# 应收账款




# 应收账款其他项目路
def addAccountsReceivableOther(document,path):
    addParagraph(document, "4、本期重要的坏账准备收回或转回情况", "paragraph")
    excelTableToWord(document, "收回或转回的坏账准备情况", path, style=2, conditions=("转回或收回金额",))

    addParagraph(document, "5、本期实际核销的应收账款情况", "paragraph")
    excelTableToWord(document, "本年实际核销的应收账款情况", path, style=2, conditions=("核销金额",))

    addParagraph(document, "6、按欠款方归集的年末余额前五名的应收账款情况", "paragraph")
    excelTableToWord(document, "按欠款方归集的年末余额前五名的应收账款情况", path, style=2, conditions=("账面余额",))

    addParagraph(document, "7、由金融资产转移而终止确认的应收账款", "paragraph")
    excelTableToWord(document, "由金融资产转移而终止确认的应收账款", path, style=2, conditions=("终止确认金额",))

    addParagraph(document, "8、转移应收账款且继续涉入形成的资产、负债", "paragraph")
    excelTableToWord(document, "转移应收账款且继续涉入形成的资产负债", path, style=2, conditions=("期末数",))

# 老准则
# 非首次执行新准则
# 本期首次执行新准则
def addAccountsReceivable(document, num, path, context):
    addTitle(document, "（{}）应收账款".format(to_chinese(num)), 2, True)
    if context["notes_params"]["tandardssForFinancialInstruments"] == "老金融工具准则":
        # 原金融工具准则
        df = filterDateFrame("应收账款期末数原金融工具准则",path,("期末账面余额",))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["种类", "期末数", "nan", "nan", "nan", "nan"],
                      ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                      ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
            addParagraph(document, "（续）", "paragraph")
        df = filterDateFrame("应收账款期初数原金融工具准则", path, ("期初账面余额",))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["种类", "期初数", "nan", "nan", "nan", "nan"],
                      ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                      ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)

        addParagraph(document, "1、期末单项金额重大并单项计提坏账准备的应收账款", "paragraph")
        excelTableToWord(document, "期末单项计提坏账准备的应收账款", path, style=2,conditions=("账面余额",))

        addParagraph(document, "2、按信用风险特征组合计提坏账准备的应收账款", "paragraph")
        addParagraph(document, "（1）采用账龄分析法计提坏账准备的应收账款", "paragraph")
        df = filterDateFrame("采用账龄分析法计提坏账准备的应收账款原准则",path,conditions=("期末账面余额","期初数账面余额"))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["账龄", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                      ["nan", "账面余额", "nan", "坏账准备", "账面余额", "nan", "坏账准备"],
                      ["nan", "金额", "比例(%)", "nan", "金额", "比例(%)", "nan"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[2, 3], [2, 6]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

        addParagraph(document, "（2）采用其他组合方法计提坏账准备的应收账款", "paragraph")
        df = filterDateFrame("采用其他组合方法计提坏账准备的应收账款原准则",path,("期末余额","期初余额"))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["组合名称", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                      ["nan", "账面余额", "计提比例（%）", "坏账准备", "账面余额", "计提比例（%）", "坏账准备"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

        addParagraph(document, "3、期末单项金额虽不重大但单项计提坏账准备的应收账款", "paragraph")
        excelTableToWord(document, "期末单项金额虽不重大但单项计提坏账准备的应收账款原准则", path, style=2,conditions=("账面余额",))
        # 添加4/5/6/7/8
        addAccountsReceivableOther(document, path)
    else:
        if "新金融工具准则" in context["standardChange"]["implementationOfNewStandardsInThisPeriod"]:
            # 首次执行新金融工具准则
            df = filterDateFrame("应收账款期末数首次新金融工具准则", path, ("期末账面余额",))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["种类", "期末数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
                addParagraph(document, "（续）", "paragraph")
            df = filterDateFrame("应收账款期初数首次新金融工具准则", path, ("期初账面余额",))
            if len(df)>0:
                titles = [["种类", "期初数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                dc = df.to_dict("split")
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")

            addParagraph(document, "1、期末单项计提坏账准备的应收账款", "paragraph")
            excelTableToWord(document, "期末单项计提坏账准备的应收账款", path, style=2,conditions=("账面余额",))

            addParagraph(document, "2、采用组合计提坏账准备的应收账款", "paragraph")
            df = filterDateFrame("采用组合计提坏账准备的应收账款首次执行",path,conditions=("期末余额",))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["组合名称", "期末数", "nan", "nan"],
                          ["nan", "账面余额", "坏账准备", "计提比例(%)"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)

                for row in dc["data"][:-1]:
                    if pd.isna(row[0]):
                        break
                    combinationName = "{}首次执行".format(row[0])
                    addParagraph(document, "{}:".format(row[0]), "paragraph")
                    df = pd.read_excel(path, sheet_name=combinationName)
                    dc = df.to_dict("split")
                    titles = [["项目", "期末数", "nan", "nan"],
                              ["nan", "账面余额", "坏账准备", "计提比例(%)"],
                              ]
                    titleLength = len(titles)
                    rowLength = len(dc["index"]) + titleLength
                    columnLength = len(dc["columns"])
                    table = createBorderedTable(document, rowLength, columnLength)
                    addCombineTitleSpecialReceivable(titles, table, [])
                    addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")

            addParagraph(document, "3、坏账准备变动明细情况", "paragraph")
            df = filterDateFrame("应收账款坏账准备变动明细情况新金融工具准则",path,conditions=("期初数","期末数"))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["项目", "期初数", "本期增加", "nan", "本期减少", "nan", "nan", "期末数"],
                          ["nan", "nan", "计提", "其他", "转回", "核销", "其他", "nan"]]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialLast(titles, table)
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")
            # 添加4/5/6/7/8
            addAccountsReceivableOther(document, path)
        else:
            # 新金融工具准则
            df = filterDateFrame("应收账款期末数新金融工具准则",path,conditions=("期末账面余额",))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["种类", "期末数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
                addParagraph(document, "（续）", "paragraph")

            df = filterDateFrame("应收账款期初数新金融工具准则", path, conditions=("期初账面余额",))
            if len(df)>0:
                titles = [["种类", "期初数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                dc = df.to_dict("split")
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)

            addParagraph(document, "1、期末单项计提坏账准备的应收账款", "paragraph")
            excelTableToWord(document, "期末单项计提坏账准备的应收账款", path, style=2,conditions=("账面余额",))

            addParagraph(document, "2、采用组合计提坏账准备的应收账款", "paragraph")
            df = filterDateFrame("采用组合计提坏账准备的应收账款新金融工具",path,conditions=("期末余额","期初余额"))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["组合名称", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                          ["nan", "账面余额", "坏账准备", "计提比例(%)", "账面余额", "坏账准备", "计提比例(%)"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)

                for row in dc["data"][:-1]:
                    combinationName = row[0]
                    addParagraph(document, "{}新金融工具:".format(row[0]), "paragraph")
                    df = pd.read_excel(path, sheet_name=combinationName)
                    dc = df.to_dict("split")
                    titles = [["项目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                              ["nan", "账面余额", "坏账准备", "计提比例(%)", "账面余额", "坏账准备", "计提比例(%)"],
                              ]
                    titleLength = len(titles)
                    rowLength = len(dc["index"]) + titleLength
                    columnLength = len(dc["columns"])
                    table = createBorderedTable(document, rowLength, columnLength)
                    addCombineTitleSpecialReceivable(titles, table, [])
                    addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")

            addParagraph(document, "3、坏账准备变动明细情况", "paragraph")
            df = filterDateFrame("应收账款坏账准备变动明细情况新金融工具准则",path,conditions=("期初数","期末数"))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["项目", "期初数", "本期增加", "nan", "本期减少", "nan", "nan", "期末数"],
                          ["nan", "nan", "计提", "其他", "转回", "核销", "其他", "nan"]]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialLast(titles, table)
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")
            # 添加4/5/6/7/8
            addAccountsReceivableOther(document, path)


def addReceivablesFinancing(document, num, path, context):
    reportDate = context["report_params"]["reportDate"]

    addTitle(document, "（{}）应收款项融资".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "应收款项融资", path, style=2)
    addParagraph(document, "本公司的应收款项融资主要为根据日常资金管理需要，预计通过转让、贴现或背书回款并终止确认的应收账款和银行承兑汇票。", "paragraph")
    addParagraph(document,
                 "本公司无单项计提减值准备的银行承兑汇票。于{}，本公司按照整个存续期预期信用损失计量坏账准备，本公司认为所持有的银行承兑汇票不存在重大信用风险，不会因银行违约而产生重大损失。".format(
                     reportDate), "paragraph")
    df = filterDateFrame("应收款项融资已转让已背书或已贴现未到期",path,conditions=("已终止确认","未终止确认"))
    if len(df)>0:
        addParagraph(document, "于{}，本公司列示于应收款项融资已转让、已背书或已贴现但尚未到期的应收票据和应收账款如下：".format(reportDate), "paragraph")
        dfToWord(document,df,style=2)


# （八）预付款项
def addAdvancePayment(document, num, path, context):
    addTitle(document, "（{}）预付款项".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、按账龄列示", "paragraph")
    df = filterDateFrame("预付账款账龄明细",path,conditions=("期末余额","期初余额"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["账  龄", "期末数", "nan", "nan", "nan", "期初数", "nan", "nan", "nan"],
                  ["nan", "账面余额", "比例(%)", "减值准备", "账面价值", "账面余额", "比例(%)", "减值准备", "账面价值"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        # addCombineTitleSpecialLast(titles, table)
        addCombineTitleSpecialReceivable(titles, table, [[1, 8]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "2、账龄超过1年的大额预付款项情况", "paragraph")
    excelTableToWord(document, "账龄超过1年的大额预付款项情况", path, style=2,conditions=("期末余额",))
    addParagraph(document, "3、按欠款方归集的年末余额前五名的预付账款情况", "paragraph")
    excelTableToWord(document, "按欠款方归集的年末余额前五名的预付账款情况", path, style=2,conditions=("账面余额",))


# 添加其他应收款披露内容
def addOtherReceivablesOthers(document,path):
    addParagraph(document, "4、本期重要的坏账准备收回或转回情况", "paragraph")
    excelTableToWord(document, "收回或转回的坏账准备情况", path, style=2, conditions=("转回或收回金额",))

    addParagraph(document, "5、本期实际核销的应收账款情况", "paragraph")
    excelTableToWord(document, "本年实际核销的应收账款情况", path, style=2, conditions=("核销金额",))

    addParagraph(document, "6、其他应收款款项性质分类情况", "paragraph")
    excelTableToWord(document, "其他应收款按性质分类情况", path, style=2)

    addParagraph(document, "7、重要逾期利息", "paragraph")
    excelTableToWord(document, "重要逾期利息", path, style=2, conditions=("期末余额",))

    addParagraph(document, "8、应收股利明细情况", "paragraph")
    excelTableToWord(document, "应收股利明细", path, style=2)

    addParagraph(document, "9、按欠款方归集的年末余额前五名的其他应收款情况", "paragraph")
    excelTableToWord(document, "按欠款方归集的年末金额前五名的其他应收款项情况", path, style=2, conditions=("账面余额",))

    addParagraph(document, "10、按应收金额确认的政府补助", "paragraph")
    excelTableToWord(document, "按应收金额确认的政府补助", path, style=2, conditions=("期末余额",))

    addParagraph(document, "11、由金融资产转移而终止确认的应收账款", "paragraph")
    excelTableToWord(document, "由金融资产转移而终止确认的应收账款", path, style=2, conditions=("终止确认金额",))

    addParagraph(document, "12、转移应收账款且继续涉入形成的资产、负债", "paragraph")
    excelTableToWord(document, "转移应收账款且继续涉入形成的资产负债", path, style=2, conditions=("期末数",))

def addOtherReceivables(document, num, path, context):
    addTitle(document, "（{}）其他应收款".format(to_chinese(num)), 2, True)
    if context["notes_params"]["tandardssForFinancialInstruments"] == "老金融工具准则":
        # 原金融工具准则
        excelTableToWord(document, "其他应收款原准则", path, style=2)
        addParagraph(document, "1、应收利息", "paragraph")
        addParagraph(document, "（1）应收利息分类", "paragraph")
        excelTableToWord(document, "应收利息分类", path, style=2)
        addParagraph(document, "（2）重要逾期利息", "paragraph")
        excelTableToWord(document, "重要逾期利息", path, style=2,conditions=("期末余额",))

        addParagraph(document, "2、应收股利", "paragraph")
        excelTableToWord(document, "应收股利明细", path, style=2)

        addParagraph(document, "3、其他应收款项", "paragraph")
        df = filterDateFrame("其他应收款项期末明细原准则",path,conditions=("期末账面余额",))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["种类", "期末数", "nan", "nan", "nan", "nan"],
                      ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                      ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
            addParagraph(document, "（续）", "paragraph")
        df = filterDateFrame("其他应收款项期初明细原准则", path, conditions=("期初账面余额",))
        if len(df)>0:
            titles = [["种类", "期初数", "nan", "nan", "nan", "nan"],
                      ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                      ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                      ]
            titleLength = len(titles)
            dc = df.to_dict("split")
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)

        addParagraph(document, "（1）年末单项金额重大并单项计提坏账准备的其他应收款项", "paragraph")
        excelTableToWord(document, "期末单项计提坏账准备的其他应收款", path, style=2,conditions=("账面余额",))

        addParagraph(document, "（2）按信用风险特征组合计提坏账准备的其他应收款项", "paragraph")
        addParagraph(document, "①采用账龄分析法计提坏账准备的其他应收款项", "paragraph")
        df = filterDateFrame("采用账龄分析法计提坏账准备的其他应收款项原准则",path,conditions=("期末账面余额","期初数账面余额"))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["账龄", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                      ["nan", "账面余额", "nan", "坏账准备", "账面余额", "nan", "坏账准备"],
                      ["nan", "金额", "比例(%)", "nan", "金额", "比例(%)", "nan"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[2, 3], [2, 6]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

        addParagraph(document, "②采用余额百分比法或其他组合方法计提坏账准备的其他应收款项", "paragraph")
        df = filterDateFrame("采用其他组合方法计提坏账准备的其他应收款原准则", path, conditions=("期末余额", "期初余额"))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["组合名称", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                      ["nan", "账面余额", "计提比例（%）", "坏账准备", "账面余额", "计提比例（%）", "坏账准备"],
                      ]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

        addParagraph(document, "（3）期末单项金额虽不重大但单项计提坏账准备的其他应收款项", "paragraph")
        excelTableToWord(document, "期末单项金额虽不重大但单项计提坏账准备的其他应收款原准则", path, style=2,conditions=("账面余额",))

        addParagraph(document, "（4）本期重要的坏账准备收回或转回情况", "paragraph")
        excelTableToWord(document, "其他应收款收回或转回的坏账准备情况", path, style=2,conditions=("转回或收回金额",))

        addParagraph(document, "（5）本期实际核销的其他应收款情况", "paragraph")
        excelTableToWord(document, "本年实际核销的其他应收款情况", path, style=2,conditions=("核销金额",))

        addParagraph(document, "（6）按欠款方归集的年末金额前五名的其他应收款项情况", "paragraph")
        excelTableToWord(document, "按欠款方归集的年末金额前五名的其他应收款项情况", path, style=2,conditions=("账面余额",))

        addParagraph(document, "（7）由金融资产转移而终止确认的其他应收款项", "paragraph")
        excelTableToWord(document, "由金融资产转移而终止确认的其他应收款项", path, style=2,conditions=("终止确认金额",))

        addParagraph(document, "（8）转移其他应收款且继续涉入形成的资产、负债", "paragraph")
        excelTableToWord(document, "转移其他应收款且继续涉入形成的资产负债", path, style=2,conditions=("期末数",))

        addParagraph(document, "（9）按应收金额确认的政府补助", "paragraph")
        excelTableToWord(document, "按应收金额确认的政府补助", path, style=2,conditions=("期末余额",))
    else:
        if "新金融工具准则" in context["standardChange"]["implementationOfNewStandardsInThisPeriod"]:
            # 首次执行新金融工具准则
            addParagraph(document, "1、明细情况", "paragraph")
            addParagraph(document, "（1）类别明细情况", "paragraph")
            df = filterDateFrame("其他应收款期末数首次新金融工具准则",path,conditions=("期末账面余额",))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["种类", "期末数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
                addParagraph(document, "（续）", "paragraph")

            df = filterDateFrame("其他应收款期初数首次新金融工具准则", path, conditions=("期初账面余额",))
            if len(df)>0:
                titles = [["种类", "期初数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                dc = df.to_dict("split")
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)

            addParagraph(document, "（2）期末单项计提坏账准备的其他应收款", "paragraph")
            excelTableToWord(document, "期末单项计提坏账准备的其他应收款", path, style=2,conditions=("账面余额",))

            addParagraph(document, "（3）采用组合计提坏账准备的其他应收款", "paragraph")
            df = filterDateFrame("采用组合计提坏账准备的其他应收款首次执行",path,conditions=("期末余额",))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["组合名称", "期末数", "nan", "nan"],
                          ["nan", "账面余额", "坏账准备", "计提比例(%)"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")

            addParagraph(document, "2、账龄情况", "paragraph")
            excelTableToWord(document, "其他应收款账龄情况新金融工具准则", path, style=2,conditions=("期末余额","期初余额"))

            addParagraph(document, "3、坏账准备变动明细情况", "paragraph")
            df = filterDateFrame("其他应收款坏账准备变动情况新金融工具准则",path,conditions=("第一阶段","第二阶段","第三阶段"))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["项目", "第一阶段", "第二阶段", "第三阶段", "合  计"],
                          ["nan", "未来12个月预期信用损失", "整个存续期预期信用损失(未发生信用减值)", "整个存续期预期信用损失(已发生信用减值)", "nan"]]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialLast(titles, table)
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")

            addOtherReceivablesOthers(document, path)
        else:
            # 新金融工具准则
            addParagraph(document, "1、明细情况", "paragraph")
            addParagraph(document, "（1）类别明细情况", "paragraph")
            df = filterDateFrame("其他应收款期末数新金融工具准则",path,conditions=("期末账面余额",))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["种类", "期末数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
                addParagraph(document, "（续）", "paragraph")
            df = filterDateFrame("其他应收款期初数新金融工具准则", path, conditions=("期初账面余额",))
            if len(df)>0:
                titles = [["种类", "期初数", "nan", "nan", "nan", "nan"],
                          ["nan", "账面余额", "nan", "坏账准备", "nan", "账面价值"],
                          ["nan", "金额", "比例(%)", "金额", "计提比例(%)", "nan"],
                          ]
                titleLength = len(titles)
                dc = df.to_dict("split")
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [[2, 5]])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)

            addParagraph(document, "（2）期末单项计提坏账准备的其他应收款", "paragraph")
            excelTableToWord(document, "期末单项计提坏账准备的其他应收款", path, style=2, conditions=("账面余额",))

            addParagraph(document, "（3）采用组合计提坏账准备的其他应收款", "paragraph")
            df = filterDateFrame("采用组合计提坏账准备的其他应收款新金融工具准则", path, conditions=("期末余额","期初余额"))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["组合名称", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                          ["nan", "账面余额", "计提比例（%）", "坏账准备", "账面余额", "计提比例（%）", "坏账准备"],
                          ]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialReceivable(titles, table, [])
                addContentToCombineTitle(document, dc, table, titleLength, style=2)

            addParagraph(document, "2、账龄情况", "paragraph")
            excelTableToWord(document, "其他应收款账龄情况新金融工具准则", path, style=2,conditions=("期末余额","期初余额"))

            addParagraph(document, "3、坏账准备变动明细情况", "paragraph")
            df = filterDateFrame("其他应收款坏账准备变动情况新金融工具准则", path, conditions=("第一阶段", "第二阶段", "第三阶段"))
            if len(df)>0:
                dc = df.to_dict("split")
                titles = [["项目", "第一阶段", "第二阶段", "第三阶段", "合  计"],
                          ["nan", "未来12个月预期信用损失", "整个存续期预期信用损失(未发生信用减值)", "整个存续期预期信用损失(已发生信用减值)", "nan"]]
                titleLength = len(titles)
                rowLength = len(dc["index"]) + titleLength
                columnLength = len(dc["columns"])
                table = createBorderedTable(document, rowLength, columnLength)
                addCombineTitleSpecialLast(titles, table)
                addContentToCombineTitle(document, dc, table, titleLength, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")

            addOtherReceivablesOthers(document, path)

# 存货
def addInventory(document, num, path, context):
    addTitle(document, "（{}）存货".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    df = filterDateFrame("存货明细情况",path,conditions=("期末余额","期初余额"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                  ["nan", "账面余额", "跌价准备", "账面价值", "账面余额", "跌价准备", "账面价值"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)


    df = filterDateFrame("房地产开发成本",path,conditions=("期末账面余额","期初账面余额"))
    if len(df)>0:
        addParagraph(document, "房地产开发成本明细如下：", "paragraph")
        dfToWord(document,df,style=2)

    df = filterDateFrame("房地产开发产品", path, conditions=("期末账面余额", "期初账面余额"))
    if len(df) > 0:
        addParagraph(document, "房地产开发产品明细如下：", "paragraph")
        dfToWord(document, df, style=2)

    df = filterDateFrame("合同履约成本", path)
    if len(df) > 0:
        addParagraph(document, "合同履约成本明细如下：", "paragraph")
        dfToWord(document, df, style=2)

    df = filterDateFrame("存货跌价准备明细情况",path)
    if len(df)>0:
        addParagraph(document, "2、存货跌价准备", "paragraph")
        dc = df.to_dict("split")
        titles = [["项  目", "期初数", "本期增加", "nan", "本期减少", "nan", "期末数"],
                  ["nan", "nan", "计提", "其他", "转回或转销", "其他", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 6]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)

        addParagraph(document, "3、确定可变现净值的具体依据、本期转回或转销存货跌价准备的原因", "paragraph")
        excelTableToWord(document, "确定可变现净值的具体依据", path, style=2,conditions=())

        df = filterDateFrame("存货期末余额中借款费用资本化情况", path)
        if len(df) > 0:
            addParagraph(document, "4、存货期末余额中借款费用资本化情况：", "paragraph")
            dfToWord(document, df, style=2)
    else:
        df = filterDateFrame("存货期末余额中借款费用资本化情况", path,conditions=("期末借款资本化余额",))
        if len(df) > 0:
            addParagraph(document, "2、存货期末余额中借款费用资本化情况：", "paragraph")
            dfToWord(document, df, style=2)


# （十一）合同资产
def addContractAssets(document, num, path, context):
    addTitle(document, "（{}）合同资产".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、合同资产情况", "paragraph")
    df = pd.read_excel(path, sheet_name="合同资产情况")
    dc = df.to_dict("split")
    titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
              ["nan", "账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值"]]
    titleLength = len(titles)
    rowLength = len(dc["index"]) + titleLength
    columnLength = len(dc["columns"])
    table = createBorderedTable(document, rowLength, columnLength)
    addCombineTableTitle(table, titles)
    addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "2、合同资产本期的重大变动", "paragraph")
    excelTableToWord(document, "合同资产本期的重大变动", path, style=2,conditions=("本期",))
    addParagraph(document, "3、合同资产减值准备", "paragraph")
    addParagraph(document, "（1）期末单项计提坏账准备的合同资产", "paragraph")
    excelTableToWord(document, "期末单项计提坏账准备的合同资产", path, style=2,conditions=("账面余额",))
    addParagraph(document, "（2）采用组合计提坏账准备的合同资产", "paragraph")
    excelTableToWord(document, "采用组合计提坏账准备的合同资产", path, style=2,conditions=("账面余额",))


# 持有待售资产
def addAssetsHeldForSale(document, num, path, context):
    addTitle(document, "（{}）持有待售资产".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、持有待售资产的基本情况", "paragraph")
    excelTableToWord(document, "持有待售资产的基本情况", path, style=2,conditions=("期末账面价值",))
    addParagraph(document, "2、持有待售资产减值准备情况", "paragraph")
    excelTableToWord(document, "持有待售资产减值准备情况", path, style=2,conditions=("期初余额","期末余额"))
    # addParagraph(document, "3、与持有待售的非流动资产或处置组有关的其他综合收益累计金额", "paragraph")
    # excelTableToWord(document, "与持有待售的非流动资产或处置组有关的其他综合收益累计金额", path, style=2,conditions=("其他综合收益累计金额",))
    # addParagraph(document, "4、本期不再划分为持有待售类别或从持有待售处置组中移除的情况", "paragraph")
    # excelTableToWord(document, "本期不再划分为持有待售类别或从持有待售处置组中移除的情况", path, style=2,conditions=("影响金额",))


# （十三）一年内到期的非流动资产
def addNonCurrentAssetsDueWithinOneYear(document, num, path, context):
    addTitle(document, "（{}）一年内到期的非流动资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "一年内到期的非流动资产", path, style=2)


# （十四）其他流动资产
def addOtherCurrentAssets(document, num, path, context):
    addTitle(document, "（{}）其他流动资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "其他流动资产", path, style=2)
    df = filterDateFrame("合同取得成本",path)
    if len(df)>0:
        addParagraph(document, "合同取得成本本期变动：", "paragraph")
        dfToWord(document,df,style=2)


# （十五）债权投资
def addDebtInvestment(document, num, path, context):
    addTitle(document, "（{}）债权投资".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    df = pd.read_excel(path, sheet_name="债权投资")
    dc = df.to_dict("split")
    titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
              ["nan", "账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值"]]
    titleLength = len(titles)
    rowLength = len(dc["index"]) + titleLength
    columnLength = len(dc["columns"])
    table = createBorderedTable(document, rowLength, columnLength)
    addCombineTableTitle(table, titles)
    addContentToCombineTitle(document, dc, table, titleLength, style=2)

    addParagraph(document, "2、债权投资减值准备", "paragraph")
    df = filterDateFrame("债权投资减值准备", path, conditions=("第一阶段", "第二阶段", "第三阶段"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项目", "第一阶段", "第二阶段", "第三阶段", "合  计"],
                  ["nan", "未来12个月预期信用损失", "整个存续期预期信用损失(未发生信用减值)", "整个存续期预期信用损失(已发生信用减值)", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialLast(titles, table)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")
    addParagraph(document, "3、期末重要的债权投资", "paragraph")
    excelTableToWord(document, "期末重要的债权投资", path, style=2,conditions=("面值",))


# （十六）可供出售金融资产
def addAvailableForSaleFinancialAssets(document, num, path, context):
    addTitle(document, "（{}）可供出售金融资产".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、可供出售金融资产情况", "paragraph")
    df = filterDateFrame("可供出售金融资产情况",path,conditions=("期末余额","期初余额"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                  ["nan", "账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    # addParagraph(document,"2、期末按成本计量的可供出售金融资产","paragraph")
    # df = pd.read_excel(path, sheet_name="期末按成本计量的可供出售金融资产")
    # dc = df.to_dict("split")
    # titles = [["被投资单位", "账面余额", "nan", "nan", "nan", "减值准备", "nan", "nan", "nan","在被投资单位持股比例（%）","本期现金红利"],
    #           ["nan", "期初", "本期增加", "本期减少","期末", "期初", "本期增加", "本期减少","期末","nan", "nan"]]
    # titleLength = len(titles)
    # rowLength = len(dc["index"]) + titleLength
    # columnLength = len(dc["columns"])
    # table = createBorderedTable(document, rowLength, columnLength)
    # addCombineTitleSpecialReceivable(titles,table,[[1,9],[1,10]])
    # addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "2、期末按公允价值计量的可供出售金融资产", "paragraph")
    excelTableToWord(document, "期末按公允价值计量的可供出售金融资产", path, style=2,conditions=("可供出售权益工具","可供出售债务工具","其他"))
    # addParagraph(document,"3、本期可供出售金融资产减值的变动情况","paragraph")
    # excelTableToWord(document,"本期可供出售金融资产减值的变动情况",path,style=2)
    addParagraph(document, "3、可供出售权益工具年末公允价值严重下跌或非暂时性下跌但未计提减值准备的相关说明", "paragraph")
    excelTableToWord(document, "可供出售权益工具严重下跌但未计提减值", path, style=2,conditions=("投资成本",))


# （十七）其他债权投资
def addOtherDebtInvestment(document, num, path, context):
    addTitle(document, "（{}）其他债权投资".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    df = filterDateFrame("其他债权投资期末数",path,conditions=("初始成本",))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "nan", "nan", "nan", "nan"],
                  ["nan", "初始成本", "利息调整", "应计利息", "公允价值变动", "账面价值", "减值准备"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
        addParagraph(document, "（续）", "paragraph")
    df = filterDateFrame("其他债权投资期初数", path, conditions=("初始成本",))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期初数", "nan", "nan", "nan", "nan", "nan"],
                  ["nan", "初始成本", "利息调整", "应计利息", "公允价值变动", "账面价值", "减值准备"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "2、其他债权投资减值准备", "paragraph")
    df = filterDateFrame("其他债权投资减值准备", path, conditions=("第一阶段", "第二阶段", "第三阶段"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项目", "第一阶段", "第二阶段", "第三阶段", "合  计"],
                  ["nan", "未来12个月预期信用损失", "整个存续期预期信用损失(未发生信用减值)", "整个存续期预期信用损失(已发生信用减值)", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialLast(titles, table)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")
    addParagraph(document, "3、期末重要的其他债权投资", "paragraph")
    excelTableToWord(document, "期末重要的其他债权投资", path, style=2)


# 持有至到期投资
def addHeldToMaturityInvestment(document, num, path, context):
    addTitle(document, "（{}）持有至到期投资".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    df = filterDateFrame("持有至到期投资明细情况",path,conditions=("期末余额","期初余额"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                  ["nan", "账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)

    addParagraph(document, "2、期末重要的持有至到期投资", "paragraph")
    excelTableToWord(document, "期末重要的持有至到期投资", path, style=2,conditions=("面值",))


# 长期应收款
def addlongTermReceivables(document, num, path, context):
    addTitle(document, "（{}）长期应收款".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    df = filterDateFrame("长期应收款明细情况",path,conditions=("期末账面余额","期初账面余额"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan", "折现率区间"],
                  ["nan", "账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "2、减值准备计提情况", "paragraph")
    df = filterDateFrame("长期应收款坏账准备变动情况新金融工具准则", path, conditions=("第一阶段", "第二阶段", "第三阶段"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项目", "第一阶段", "第二阶段", "第三阶段", "合  计"],
                  ["nan", "未来12个月预期信用损失", "整个存续期预期信用损失(未发生信用减值)", "整个存续期预期信用损失(已发生信用减值)", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialLast(titles, table)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")
    addParagraph(document, "3、因金融资产转移而终止确认的长期应收款", "paragraph")
    excelTableToWord(document, "因金融资产转移而终止确认的长期应收款", path, style=2,conditions=("终止确认的长期应收款金额",))
    addParagraph(document, "4、转移长期应收款且继续涉入形成的资产、负债金额", "paragraph")
    excelTableToWord(document, "转移长期应收款且继续涉入形成的资产负债金额", path, style=2,conditions=("期末数",))


def addLongTermEquityInvestmentOther(document,df):
    dc = df.to_dict("split")
    titles = [["被投资单位", "期初余额", "本期增减变动", "nan", "nan ", "nan", "nan", "nan ", "nan", "nan", "期末余额", "减值准备期末余额"],
              ["nan", "nan", "追加投资", "减少投资", "权益法下确认的投资损益", "其他综合收益调整", "其他权益变动", "宣告发放现金股利或利润", "计提减值准备", "其他", "nan",
               "nan"]]
    titleLength = len(titles)
    rowLength = len(dc["index"]) + titleLength
    columnLength = len(dc["columns"])
    table = createBorderedTable(document, rowLength, columnLength)
    addCombineTitleSpecialReceivable(titles, table, [[1, 10], [1, 11]])
    addContentToCombineTitle(document, dc, table, titleLength, style=2)

# （二十）长期股权投资
def addContent1(document, reportType, startYear, path):
    addParagraph(document, "2、明细情况", "paragraph")
    df = filterDateFrame("长期股权投资子公司明细情况", path)
    if len(df)>0:
        addParagraph(document, "（1）子公司", "paragraph")
        dfToWord(document,df,style=2)
        df = filterDateFrame("长期股权投资合营企业明细情况", path)
        if len(df)>0:
            addParagraph(document, "（2）合营企业", "paragraph")
            addLongTermEquityInvestmentOther(document, df)
            df = filterDateFrame("长期股权投资联营企业明细情况", path)
            if len(df)>0:
                addParagraph(document, "（3）联营企业", "paragraph")
                addLongTermEquityInvestmentOther(document, df)
            else:
                addParagraph(document, "不适用", "paragraph")
        else:
            if len(df)>0:
                addParagraph(document, "（2）联营企业", "paragraph")
                addLongTermEquityInvestmentOther(document, df)
            else:
                addParagraph(document, "不适用", "paragraph")
    else:
        df = filterDateFrame("长期股权投资合营企业明细情况", path)
        if len(df)>0:
            addParagraph(document, "（1）合营企业", "paragraph")
            addLongTermEquityInvestmentOther(document, df)
            df = filterDateFrame("长期股权投资联营企业明细情况", path)
            if len(df)>0:
                addParagraph(document, "（2）联营企业", "paragraph")
                addLongTermEquityInvestmentOther(document, df)
            else:
                addParagraph(document, "不适用", "paragraph")
        else:
            df = filterDateFrame("长期股权投资联营企业明细情况", path)
            if len(df)>0:
                addParagraph(document, "（1）联营企业", "paragraph")
                addLongTermEquityInvestmentOther(document, df)
            else:
                addParagraph(document, "不适用", "paragraph")



def addLongTermEquityInvestment(document, num, path, context):
    companyType = context["report_params"]["companyType"]
    # 报告日期
    reportDate = context["report_params"]["reportDate"]
    # 获取报告起始日
    startYear = reportDate[:4]

    addTitle(document, "（{}）长期股权投资".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、分类情况", "paragraph")
    df = filterDateFrame("长期股权投资分类情况",path)
    if len(df)>0:
        dfToWord(document,df,style=2)

    addLandscapeContent(document, addContent1, type, startYear, path)

    if companyType == "国有企业":
        addParagraph(document, "3、对合营企业投资和联营企业投资", "paragraph")
        excelTableToWord(document, "对合营企业投资和联营企业投资国有企业", path, style=2,conditions=("期末资产总额",))

        addParagraph(document, "4、向投资企业转移资金的能力受到限制的有关情况", "paragraph")
        df = pd.read_excel(path, sheet_name="向投资企业转移资金的能力受到限制的有关情况国有企业")
        if len(df)>0:
            dfToWord(document,df,style=2)
        else:
            addParagraph(document, "不适用", "paragraph")

    # addParagraph(document,"3、重要合营企业的主要财务信息","paragraph")
    # addParagraph(document,"本期数：","paragraph")
    # excelTableToWord(document,"重要合营企业财务信息本期数",path,style=2)
    # addParagraph(document,"上期数：","paragraph")
    # excelTableToWord(document,"重要合营企业财务信息上期数",path,style=2)
    #
    # addParagraph(document,"4、重要联营企业的主要财务信息","paragraph")
    # addParagraph(document,"本期数：","paragraph")
    # excelTableToWord(document,"重要联营企业财务信息本期数",path,style=2)
    # addParagraph(document,"上期数：","paragraph")
    # excelTableToWord(document,"重要联营企业财务信息上期数",path,style=2)
    #
    # addParagraph(document,"5、重要联营企业的主要财务信息","paragraph")
    # excelTableToWord(document,"不重要合营企业和联营企业的汇总信息",path,style=2)
    #
    # addParagraph(document,"6、合营企业或联营企业发生的超额亏损","paragraph")
    # excelTableToWord(document,"合营企业或联营企业发生的超额亏损",path,style=2)


def addInvestmentInOtherEquityInstruments(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）其他权益工具投资".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    excelTableToWord(document, "其他权益工具投资明细", path, style=2)
    if companyType == "上市公司":
        addParagraph(document, "2、非交易性权益工具投资情况", "paragraph")
        excelTableToWord(document, "非交易性权益工具投资情况上市公司", path, style=2,conditions=("累计利得","累计损失"))
    else:
        addParagraph(document, "2、期末重要的其他权益工具投资", "paragraph")
        excelTableToWord(document, "期末重要的其他权益工具投资国有企业", path, style=2,conditions=("投资成本",))


def addOtherNonCurrentFinancialAssets(document, num, path, context):
    addTitle(document, "（{}）其他非流动金融资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "其他非流动金融资产", path, style=2)


def addInvestmentRealEstate(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）投资性房地产".format(to_chinese(num)), 2, True)
    if context["notes_params"]["measurementModeOfInvestmentRealEstate"] == "成本模式":
        if companyType == "国有企业":
            addParagraph(document, "1、采用成本计量模式的投资性房地产", "paragraph")
            excelTableToWord(document, "采用成本计量模式的投资性房地产国有企业", path, style=3)
        else:
            addParagraph(document, "1、采用成本计量模式的投资性房地产", "paragraph")
            excelTableToWord(document, "采用成本计量模式的投资性房地产上市公司", path, style=3,conditions=("合计",))
    else:
        addParagraph(document, "1、采用公允价值计量模式的投资性房地产", "paragraph")
        excelTableToWord(document, "采用公允价值计量模式的投资性房地产", path, style=3,conditions=("合计",))
    addParagraph(document, "2、未办妥产权证书的投资性房地产金额及原因", "paragraph")
    excelTableToWord(document, "未办妥产权证书的投资性房地产金额及原因", path, style=2,conditions=("账面价值",))


def addFixedAssets(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）固定资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "固定资产汇总", path, style=2,conditions=())
    addParagraph(document, "1、固定资产", "paragraph")
    addParagraph(document, "（1）固定资产情况", "paragraph")
    if companyType == "上市公司":
        excelTableToWord(document, "固定资产情况上市公司", path, style=3,conditions=("合计",))
    else:
        excelTableToWord(document, "固定资产情况国有企业", path, style=3)
    addParagraph(document, "（2）暂时闲置的固定资产情况", "paragraph")
    excelTableToWord(document, "暂时闲置的固定资产情况", path, style=2,conditions=("账面原值",))
    addParagraph(document, "（3）通过经营租赁租出的固定资产", "paragraph")
    excelTableToWord(document, "通过经营租赁租出的固定资产", path, style=2,conditions=("期末账面价值",))
    addParagraph(document, "（4）未办妥产权证书的固定资产情况", "paragraph")
    excelTableToWord(document, "未办妥产权证书的固定资产情况", path, style=2,conditions=("账面价值",))
    addParagraph(document, "2、固定资产清理", "paragraph")
    excelTableToWord(document, "固定资产清理", path, style=2)


# 在建工程
def addContent2(document, reportType, startYear, path):
    addParagraph(document, "（2）重要在建工程项目本期变动情况", "paragraph")
    excelTableToWord(document, "重要在建工程项目本期变动情况", path, style=2)


def addConstructionInProgress(document, num, path, context):
    # 报告日期
    reportDate = context["report_params"]["reportDate"]
    # 获取报告起始日
    startYear = reportDate[:4]

    addTitle(document, "（{}）在建工程".format(to_chinese(num)), 2, True)
    df = pd.read_excel(path, sheet_name="在建工程汇总")
    dfToWord(document,df,style=2)
    addParagraph(document, "1、在建工程", "paragraph")
    addParagraph(document, "（1）在建工程情况", "paragraph")
    df = filterDateFrame("在建工程情况",path,conditions=("期末余额","期初余额"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "nan", "期初数", "nan", "nan"],
                  ["nan", "账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")

    addLandscapeContent(document, addContent2, type, startYear, path)

    addParagraph(document, "（3）本期计提在建工程减值准备情况", "paragraph")
    excelTableToWord(document, "本期计提在建工程减值准备情况", path, style=2,conditions=("本期计提金额",))

    addParagraph(document, "2、工程物资", "paragraph")
    excelTableToWord(document, "工程物资", path, style=2)


def addProductiveBiologicalAssets(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）生产性生物资产".format(to_chinese(num)), 2, True)
    if companyType == "上市公司":
        excelTableToWord(document, "生产性生物资产上市公司", path, style=3,conditions=("合计",))
    else:
        excelTableToWord(document, "生产性生物资产国有企业", path, style=3)


def addOilAndGasAssets(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）油气资产".format(to_chinese(num)), 2, True)
    if companyType == "上市公司":
        excelTableToWord(document, "油气资产上市公司", path, style=3,conditions=("合计",))
    else:
        excelTableToWord(document, "油气资产国有企业", path, style=3)


def addRightToUseAssets(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）使用权资产".format(to_chinese(num)), 2, True)
    if companyType == "上市公司":
        excelTableToWord(document, "使用权资产上市公司", path, style=3,conditions=("合计",))
    else:
        excelTableToWord(document, "使用权资产国有企业", path, style=3)


def addIntangibleAssets(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "（{}）无形资产".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    if companyType == "上市公司":
        excelTableToWord(document, "无形资产上市公司", path, style=3,conditions=("合计",))
    else:
        excelTableToWord(document, "无形资产国有企业", path, style=3)
    addParagraph(document, "2、未办妥产权证书的土地使用权情况", "paragraph")
    excelTableToWord(document, "未办妥产权证书的土地使用权情况", path, style=2,conditions=("账面价值",))


def addDevelopmentExpenditure(document, num, path, context):
    addTitle(document, "（{}）开发支出".format(to_chinese(num)), 2, True)
    df = filterDateFrame("开发支出",path)
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期初数", "本期增加金额", "nan", "本期减少金额", "nan", "nan", "期末数"],
                  ["nan", "nan", "内部开发支出", "其他", "确认为无形资产", "转入当期损益", "其他", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)


def addGoodwill(document, num, path, context):
    addTitle(document, "（{}）商誉".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、商誉账面价值", "paragraph")
    excelTableToWord(document, "商誉账面价值", path, style=2)
    addParagraph(document, "2、商誉减值准备", "paragraph")
    excelTableToWord(document, "商誉减值准备", path, style=2)


def addLongTermDeferredExpenses(document, num, path, context):
    addTitle(document, "（{}）长期待摊费用".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "长期待摊费用", path, style=2)


def addDeferredTax(document, num, path, context):
    addTitle(document, "（{}）递延所得税资产和递延所得税负债".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、未经抵销的递延所得税资产", "paragraph")
    df = filterDateFrame("未经抵销的递延所得税资产",path,conditions=("期末递延所得税资产","期初递延所得税资产"))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "期初数", "nan"],
                  ["nan", "可抵扣暂时性差异", "递延所得税资产", "可抵扣暂时性差异", "递延所得税资产"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")

    df = filterDateFrame("未经抵销的递延所得税资产", path, conditions=("期末递延所得税负债", "期初递延所得税负债"))
    if len(df)>0:
        addParagraph(document, "2、未经抵销的递延所得税负债", "paragraph")
        df = pd.read_excel(path, sheet_name="未经抵销的递延所得税负债")
        dc = df.to_dict("split")
        titles = [["项  目", "期末数", "nan", "期初数", "nan"],
                  ["nan", "应纳税暂时性差异", "递延所得税负债", "应纳税暂时性差异", "递延所得税负债"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")
    addParagraph(document, "3、未确认递延所得税资产明细", "paragraph")
    excelTableToWord(document, "未确认递延所得税资产明细", path, style=2)
    addParagraph(document, "4、未确认递延所得税资产的可抵扣亏损将于以下年度到期", "paragraph")
    excelTableToWord(document, "未确认递延所得税资产的可抵扣亏损将于以下年度到期", path, style=2)


# 其他非流动资产
def addOtherNonCurrentAssets(document, num, path, context):
    addTitle(document, "（{}）其他非流动资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "其他非流动资产", path, style=2,sort=True,sort_column_name="期末数")


# 短期借款
def addShortTermLoan(document, num, path, context):
    addTitle(document, "（{}）短期借款".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、	明细情况", "paragraph")
    excelTableToWord(document, "短期借款明细情况", path, style=2,sort=True,sort_column_name="期末数")
    addParagraph(document, "2、已逾期未偿还的短期借款情况", "paragraph")
    excelTableToWord(document, "已逾期未偿还的短期借款情况", path, style=2,conditions=("期末数",))


# 交易性金融负债
def addTradingFinancialLiabilities(document, num, path, context):
    addTitle(document, "（{}）交易性金融负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "交易性金融负债", path, style=2,sort=True,sort_column_name="期末数")


# 以公允价值计量且其变动计入当期损益的金融负债
def addFinancialLiabilitiesMeasuredAtFairValueWithChangesIncludedInCurrentProfitAndLoss(document, num, path, context):
    addTitle(document, "（{}）以公允价值计量且其变动计入当期损益的金融负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "以公允价值计量且其变动计入当期损益的金融负债", path, style=2)


def addDerivativeFinancialLiabilities(document, num, path, context):
    addTitle(document, "（{}）衍生金融负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "衍生金融负债", path, style=2)


def addNotesPayable(document, num, path, context):
    addTitle(document, "（{}）应付票据".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "应付票据", path, style=2)


def addAccountsPayable(document, num, path, context):
    addTitle(document, "（{}）应付账款".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "应付账款", path, style=2)
    df = filterDateFrame("账龄超过一年的重要应付账款",path,conditions=("期末数",))
    if len(df)>0:
        addParagraph(document, "其中：账龄超过1年的重要应付账款", "paragraph")
        dfToWord(document,df,style=2)


def addDepositReceived(document, num, path, context):
    addTitle(document, "（{}）预收款项".format(to_chinese(num)), 2, True)
    companyType = context["report_params"]["companyType"]
    if companyType=="国有企业":
        excelTableToWord(document, "预收款项账龄表", path, style=2)
    else:
        excelTableToWord(document, "预收款项", path, style=2)
    df = filterDateFrame("账龄一年以上重要的预收款项", path, conditions=("期末数",))
    if len(df) > 0:
        addParagraph(document, "其中：账龄超过1年的重要预收款项", "paragraph")
        dfToWord(document, df, style=2)


def addContractualLiabilities(document, num, path, context):
    addTitle(document, "（{}）合同负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "合同负债", path, style=2)


def addEmployeeCompensationPayable(document, num, path, context):
    addTitle(document, "（{}）应付职工薪酬".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    excelTableToWord(document, "应付职工薪酬明细情况", path, style=2,conditions=("期初数","本期增加","本期减少","期末数"))
    addParagraph(document, "2、短期薪酬列示", "paragraph")
    excelTableToWord(document, "短期薪酬列示", path, style=2,conditions=("期初数","本期增加","本期减少","期末数"))
    addParagraph(document, "3、设定提存计划列示", "paragraph")
    excelTableToWord(document, "设定提存计划列示", path, style=2,conditions=("期初数","本期增加","本期减少","期末数"))


def addTaxPayable(document, num, path, context):
    addTitle(document, "（{}）应交税费".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "应交税费", path, style=2,conditions=("期初数","本期应交","本期已交","期末数"),sort=True,sort_column_name="期末数")


def addOtherAccountsPayable(document, num, path, context):
    addTitle(document, "（{}）其他应付款".format(to_chinese(num)), 2, True)
    df = pd.read_excel(path, sheet_name="其他应付款汇总")
    dfToWord(document,df,style=2)
    addParagraph(document, "1、应付利息", "paragraph")
    excelTableToWord(document, "应付利息", path, style=2)
    df = filterDateFrame("重要的已逾期未支付的利息情况",path,conditions=("逾期金额",))
    if len(df)>0:
        addParagraph(document, "其中：重要的已逾期未支付的利息情况", "paragraph")
        dfToWord(document,path,style=2)

    addParagraph(document, "2、应付股利", "paragraph")
    excelTableToWord(document, "应付股利", path, style=2)
    df = filterDateFrame("账龄一年以上重要的应付股利",path,conditions=("未支付金额",))
    if len(df)>0:
        addParagraph(document, "其中：账龄1年以上重要的应付股利", "paragraph")
        dfToWord(document,path,style=2)
    addParagraph(document, "3、其他应付款项", "paragraph")
    excelTableToWord(document, "其他应付款项", path, style=2)
    df = filterDateFrame("账龄超过一年的重要其他应付款项",path,conditions=("期末余额",))
    if len(df)>0:
        addParagraph(document, "其中：账龄超过1年的重要其他应付款项", "paragraph")
        dfToWord(document,df,style=2)


def addLiabilitiesHeldForSale(document, num, path, context):
    addTitle(document, "（{}）持有待售负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "持有待售负债", path, style=2)


def addNonCurrentLiabilitiesDueWithinOneYear(document, num, path, context):
    addTitle(document, "（{}）一年内到期的非流动负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "一年内到期的非流动负债", path, style=2,sort=True,sort_column_name="期末数")


# 其他流动负债
def addContent3(document, reportType, startYear, path):
    df = filterDateFrame("短期应付债券",path,conditions=("面值",))
    addParagraph(document, "其中：短期应付债券的增减变动", "paragraph")
    dfToWord(document,df,style=2)


def addOtherCurrentLiabilities(document, num, path, context):
    # 报告日期
    reportDate = context["report_params"]["reportDate"]
    # 获取报告起始日
    startYear = reportDate[:4]

    addTitle(document, "（{}）其他流动负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "其他流动负债", path, style=2,sort=True,sort_column_name="期末数")
    df = filterDateFrame("短期应付债券", path, conditions=("面值",))
    if len(df) > 0:
        addLandscapeContent(document, addContent3, type, startYear, path)


# 长期借款
def addLongTermLoan(document, num, path, context):
    addTitle(document, "（{}）长期借款".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "长期借款", path, style=2,sort=True,sort_column_name="期末数")


def addBondsPayable(document, num, path, context):
    # 报告日期
    reportDate = context["report_params"]["reportDate"]
    # 获取报告起始日
    startYear = reportDate[:4]

    addTitle(document, "（{}）应付债券".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    excelTableToWord(document, "应付债券", path, style=2,sort=True,sort_column_name="期末数")

    def addContent4(document, reportType, startYear, path):
        df = filterDateFrame("应付债券的增减变动",path,conditions=("面值",))
        df = df_sort(df,index_name="债券名称",column_name="期末数")
        addParagraph(document, "2、应付债券增减变动（不包括划分为金融负债的优先股、永续债等其他金融工具）", "paragraph")
        dfToWord(document,df,style=2)

    df = filterDateFrame("应付债券的增减变动", path, conditions=("面值",))
    if len(df) > 0:
        addLandscapeContent(document, addContent4, type, startYear, path)


# 优先股、永续债等金融工具
def addPreferredSharesAndPerpetualBonds(document, num, path, context):
    addTitle(document, "（{}）优先股、永续债等金融工具".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、	期末发行在外的优先股、永续债等金融工具情况", "paragraph")
    excelTableToWord(document, "期末发行在外的优先股永续债等金融工具情况", path, style=2)
    addParagraph(document, "2、	发行在外的优先股、永续债等金融工具变动情况", "paragraph")
    df = pd.read_excel(path, sheet_name="发行在外的优先股永续债等金融工具变动情况")
    dc = df.to_dict("split")
    titles = [["发行在外的金融工具", "期初余额", "nan", "本期增加", "nan", "本期减少", "nan", "期末余额", "nan"],
              ["nan", "数量", "账面价值", "数量", "账面价值", "数量", "账面价值", "数量", "账面价值"]]
    titleLength = len(titles)
    rowLength = len(dc["index"]) + titleLength
    columnLength = len(dc["columns"])
    table = createBorderedTable(document, rowLength, columnLength)
    addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
    addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "3、	归属于权益工具持有者的信息", "paragraph")
    excelTableToWord(document, "归属于权益工具持有者的信息", path, style=2)


def addLeaseLiabilities(document, num, path, context):
    addTitle(document, "（{}）租赁负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "租赁负债", path, style=3)


def addLongTermAccountsPayable(document, num, path, context):
    addTitle(document, "（{}）长期应付款".format(to_chinese(num)), 2, True)
    df = pd.read_excel(path, sheet_name="长期应付款汇总")
    dfToWord(document,df,style=2)
    addParagraph(document, "1、长期应付款", "paragraph")
    excelTableToWord(document, "长期应付款", path, style=2,sort=True,sort_column_name="期末数")
    addParagraph(document, "2、专项应付款", "paragraph")
    excelTableToWord(document, "专项应付款", path, style=2,sort=True,sort_column_name="期末数")


def addLongTermEmployeeCompensationPayable(document, num, path, context):
    addTitle(document, "（{}）长期应付职工薪酬".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    excelTableToWord(document, "长期应付职工薪酬明细情况", path, style=2)
    addParagraph(document, "2、设定受益计划变动情况", "paragraph")
    addParagraph(document, "（1）设定受益计划义务现值", "paragraph")
    excelTableToWord(document, "设定受益计划义务现值", path, style=2,conditions=("本期数","上期数"))
    addParagraph(document, "（2）计划资产", "paragraph")
    excelTableToWord(document, "计划资产", path, style=2,conditions=("本期数","上期数"))
    addParagraph(document, "（3）设定受益计划净负债（净资产）", "paragraph")
    excelTableToWord(document, "设定受益计划净负债", path, style=2,conditions=("本期数","上期数"))


def addEstimatedLiabilities(document, num, path, context):
    addTitle(document, "（{}）预计负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "预计负债", path, style=2,sort=True,sort_column_name="期末数")


def addDeferredIncome(document, num, path, context):
    addTitle(document, "（{}）递延收益".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "递延收益", path, style=2)
    df = filterDateFrame("递延收益中政府补助项目",path)
    if len(df)>0:
        addParagraph(document, "其中，涉及政府补助的项目：", "paragraph")
        dfToWord(document,df,style=2)


def addOtherNonCurrentLiabilities(document, num, path, context):
    addTitle(document, "（{}）其他非流动负债".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "其他非流动负债", path, style=2,sort=True,sort_column_name="期末数")


def addEquity(document, num, path, context):
    companyType = context["report_params"]["companyType"]

    if companyType == "国有企业":
        addTitle(document, "（{}）实收资本".format(to_chinese(num)), 2, True)
        df = filterDateFrame("实收资本",path,conditions=("期初投资金额","本期增加","本期减少","期末投资金额"))
        if len(df)>0:
            dc = df.to_dict("split")
            titles = [["投资者名称", "期初余额", "nan", "本期增加", "本期减少", "期末余额", "nan"],
                      ["nan", "投资金额", "所占比例(%)", "nan", "nan", "投资金额", "所占比例(%)"]]
            titleLength = len(titles)
            rowLength = len(dc["index"]) + titleLength
            columnLength = len(dc["columns"])
            table = createBorderedTable(document, rowLength, columnLength)
            addCombineTitleSpecialReceivable(titles, table, [[1, 3], [1, 4]])
            addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addTitle(document, "（{}）股本".format(to_chinese(num)), 2, True)
        df = pd.read_excel(path, sheet_name="股本")
        dc = df.to_dict("split")
        titles = [["项  目", "期初数", "本期增减变动（减少以“—”表示）", "nan", "nan", "nan", "nan", "期末数"],
                  ["nan", "nan", "发行新股", "送股", "公积金转股", "其他", "小计", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)


def addOtherEquityInstruments(document, num, path, context):
    addTitle(document, "（{}）其他权益工具".format(to_chinese(num)), 2, True)
    df = filterDateFrame("其他权益工具",path,conditions=("期初账面价值","期末账面价值"))
    dc = df.to_dict("split")
    titles = [["项  目", "期初余额", "nan", "本期增加", "nan", "本期减少", "nan", "期末余额", "nan"],
              ["nan", "数量", "账面价值", "数量", "账面价值", "数量", "账面价值", "数量", "账面价值"]]
    titleLength = len(titles)
    rowLength = len(dc["index"]) + titleLength
    columnLength = len(dc["columns"])
    table = createBorderedTable(document, rowLength, columnLength)
    addCombineTitleSpecialReceivable(titles, table, [[1, 7]])
    addContentToCombineTitle(document, dc, table, titleLength, style=2)


def addCapitalReserve(document, num, path, context):
    addTitle(document, "（{}）资本公积".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "资本公积", path, style=2)


def addOtherComprehensiveIncome(document, num, path, context):
    addTitle(document, "（{}）其他综合收益".format(to_chinese(num)), 2, True)
    df = filterDateFrame("其他综合收益",path)
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "期初数", "本期发生额", "nan", "nan", "nan", "nan", "nan", "期末数"],
                  ["nan", "nan", "本期所得税前发生额", "减：前期计入其他综合收益当期转入损益", "减：前期计入其他综合收益当期转入留存收益", "减：所得税费用", "税后归属于母公司",
                   "税后归属于少数股东", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 8]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")


def addSpecialReserve(document, num, path, context):
    addTitle(document, "（{}）专项储备".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "专项储备", path, style=2)


def addSurplusReserve(document, num, path, context):
    addTitle(document, "（{}）盈余公积".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "盈余公积", path, style=2)


def addUndistributedProfit(document, num, path, context):
    addTitle(document, "（{}）未分配利润".format(to_chinese(num)), 2, True)
    # df = pd.read_excel(path, sheet_name="未分配利润")
    # dfToWord(document, df, style=3)
    excelTableToWord(document, "未分配利润", path, style=3,conditions=("本期数","上年同期数"))


def addRevenueCost(document, num, path, context):
    addTitle(document, "（{}）营业收入、营业成本".format(to_chinese(num)), 2, True)
    df = pd.read_excel(path, sheet_name="营业收入与营业成本")
    dc = df.to_dict("split")
    titles = [["项  目", "本期数", "nan", "上年同期数", "nan"],
              ["nan", "收入", "成本", "收入", "成本"]]
    titleLength = len(titles)
    rowLength = len(dc["index"]) + titleLength
    columnLength = len(dc["columns"])
    table = createBorderedTable(document, rowLength, columnLength)
    addCombineTableTitle(table, titles)
    addContentToCombineTitle(document, dc, table, titleLength, style=2)
    addParagraph(document, "1、主营业务收入和主营业务成本", "paragraph")
    df = filterDateFrame("主营业务收入与主营业务成本",path,conditions=("本期收入","本期成本","上年同期收入","上年同期成本"))
    df = df_sort(df,column_name="本期收入")
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "本期数", "nan", "上年同期数", "nan"],
                  ["nan", "收入", "成本", "收入", "成本"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")
    addParagraph(document, "2、其他业务收入和其他业务成本", "paragraph")
    df = filterDateFrame("其他业务收入与其他业务成本", path, conditions=("本期收入", "本期成本", "上年同期收入", "上年同期成本"))
    df = df_sort(df, column_name="本期收入")
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["项  目", "本期数", "nan", "上年同期数", "nan"],
                  ["nan", "收入", "成本", "收入", "成本"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")


def addTaxesAndSurcharges(document, num, path, context):
    addTitle(document, "（{}）税金及附加".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "税金及附加", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addsellingExpenses(document, num, path, context):
    addTitle(document, "（{}）销售费用".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "销售费用", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addAdministrativeExpenses(document, num, path, context):
    addTitle(document, "（{}）管理费用".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "管理费用", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addRDExpenses(document, num, path, context):
    addTitle(document, "（{}）研发费用".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "研发费用", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addFinancialExpenses(document, num, path, context):
    addTitle(document, "（{}）财务费用".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "财务费用", path, style=2,conditions=("本期数","上年同期数"))


def addOtherIncome(document, num, path, context):
    addTitle(document, "（{}）其他收益".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "其他收益", path, style=2,conditions=("本期数","上年同期数"))


def addincomeFromInvestment(document, num, path, context):
    addTitle(document, "（{}）投资收益".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "投资收益", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addNetExposureHedgingGains(document, num, path, context):
    addTitle(document, "（{}）净敞口套期收益".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "净敞口套期收益", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addIncomeFromChangesInFairValue(document, num, path, context):
    addTitle(document, "（{}）公允价值变动收益".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "公允价值变动损益", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addCreditImpairmentLoss(document, num, path, context):
    addTitle(document, "（{}）信用减值损失".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "信用减值损失", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addassetsImpairmentLoss(document, num, path, context):
    addTitle(document, "（{}）资产减值损失".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "资产减值损失", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addIncomeFromAssetDisposal(document, num, path, context):
    addTitle(document, "（{}）资产处置收益".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "资产处置收益", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addNonOperatingIncome(document, num, path, context):
    addTitle(document, "（{}）营业外收入".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "营业外收入", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addNonOperatingExpenses(document, num, path, context):
    addTitle(document, "（{}）营业外支出".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "营业外支出", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addIncomeTaxExpenses(document, num, path, context):
    addTitle(document, "（{}）所得税费用".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、明细情况", "paragraph")
    excelTableToWord(document, "所得税费用", path, style=2,conditions=("本期数","上年同期数"))
    addParagraph(document, "2、会计利润与所得税费用调整过程", "paragraph")
    excelTableToWord(document, "会计利润与所得税费用调整过程", path, style=3,conditions=("本期数","上年同期数"))


def addOtherCashReceivedRelatedToOperatingActivities(document, num, path, context):
    addTitle(document, "（{}）收到其他与经营活动有关的现金".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "收到其他与经营活动有关的现金", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addOtherCashPaymentsRelatedToOperatingActivities(document, num, path, context):
    addTitle(document, "（{}）支付其他与经营活动有关的现金".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "支付其他与经营活动有关的现金", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addOtherCashReceivedRelatedToInvestmentActivities(document, num, path, context):
    addTitle(document, "（{}）收到其他与投资活动有关的现金".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "收到其他与投资活动有关的现金", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addOtherCashPaymentsRelatedToInvestmentActivities(document, num, path, context):
    addTitle(document, "（{}）支付其他与投资活动有关的现金".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "支付其他与投资活动有关的现金", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addOtherCashReceivedRelatedToFinancingActivities(document, num, path, context):
    addTitle(document, "（{}）收到其他与筹资活动有关的现金".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "收到其他与筹资活动有关的现金", path, style=2,conditions=("本期数","上年同期数"),sort=True)


def addOtherCashPaymentsRelatedToFinancingActivities(document, num, path, context):
    addTitle(document, "（{}）支付其他与筹资活动有关的现金".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "支付其他与筹资活动有关的现金", path, style=2,conditions=("本期数","上年同期数"),sort=True)


# addTitle(document, "（七十七）每股收益", 2, True)
# addTitle(document, "（七十八）非货币性资产交换", 2, True)
# addTitle(document, "（七十九）股份支付", 2, True)
def addShareBasedPayment(document, path, context):
    companyType = context["report_params"]["companyType"]
    if companyType == "上市公司":
        addTitle(document, "(一) 股份支付总体情况", 2, True)
        df = pd.read_excel(path, sheet_name="股份支付总体情况")
        dfToWord(document,df,style=2)
        addTitle(document, "(二) 以权益结算的股份支付情况", 2, True)
        df = pd.read_excel(path, sheet_name="以权益结算的股份支付情况")
        dfToWord(document, df, style=2)
        addTitle(document, "(三) 以现金结算的股份支付情况", 2, True)
        df = pd.read_excel(path, sheet_name="以现金结算的股份支付情况")
        dfToWord(document, df, style=2)
    else:
        addTitle(document, "1、股份支付总体情况", 2, True)
        df = pd.read_excel(path, sheet_name="股份支付总体情况")
        dfToWord(document, df, style=2)
        addTitle(document, "2、以权益结算的股份支付情况", 2, True)
        df = pd.read_excel(path, sheet_name="以权益结算的股份支付情况")
        dfToWord(document, df, style=2)
        addTitle(document, "3、以现金结算的股份支付情况", 2, True)
        df = pd.read_excel(path, sheet_name="以现金结算的股份支付情况")
        dfToWord(document, df, style=2)


# addTitle(document, "（八十）债务重组", 2, True)
# addTitle(document, "（八十一）借款费用", 2, True)
# addTitle(document, "（八十二）外币折算", 2, True)
# addTitle(document, "（八十三）租赁", 2, True)
# addTitle(document, "（八十四）终止经营", 2, True)
# addTitle(document, "（八十五）分部信息", 2, True)
# addTitle(document, "（八十六）合并现金流量表", 2, True)


def addCashFlowReplenishment(document, num, path, context):
    # 报告类型
    reportType = context["report_params"]["type"]
    # 现金流量表补充资料
    addTitle(document, "（{}）现金流量表补充资料".format(to_chinese(num)), 2, True)
    addParagraph(document, "1、现金流量表补充资料", "paragraph")
    try:
        addParagraph(document, "将净利润调节为经营活动现金流量", "paragraph")
        excelTableToWord(document, "将净利润调节为经营活动现金流量", path, style=2, conditions=("本期数","上年同期数"))
        addParagraph(document, "不涉及现金收支的重大投资和筹资活动", "paragraph")
        excelTableToWord(document, "不涉及现金收支的重大投资和筹资活动", path, style=2, conditions=("本期数", "上年同期数"))
        addParagraph(document, "现金及现金等价物净变动情况", "paragraph")
        excelTableToWord(document, "现金及现金等价物净变动情况", path, style=2, conditions=("本期数", "上年同期数"))
    except Exception as e:
        excelTableToWord(document, "现金流量表补充资料", path, style=2, conditions=("本期数", "上年同期数"))

    # df1 = pd.read_excel(path, sheet_name="现金及现金等价物的构成")
    df1 = filterDateFrame("现金及现金等价物的构成",path)
    if reportType == "合并":
        if context["noteAppend"]["purchaseSubsidiaries"]:
            df = filterDateFrame("本期支付的取得子公司的现金净额",path,conditions=("本期数",))
            if len(df)>0:
                addParagraph(document, "2、本期支付的取得子公司的现金净额", "paragraph")
                dfToWord(document,df,style=2)
            if context["noteAppend"]["disposalSubsidiaries"]:
                df = filterDateFrame("本期收到的处置子公司的现金净额", path, conditions=("本期数",))
                if len(df) > 0:
                    addParagraph(document, "3、本期收到的处置子公司的现金净额", "paragraph")
                    dfToWord(document, df, style=2)
            else:
                addParagraph(document, "3、现金及现金等价物的构成", "paragraph")
                dfToWord(document, df1, style=2)
                # excelTableToWord(document, "现金及现金等价物的构成", path, style=2)
        else:
            if context["noteAppend"]["disposalSubsidiaries"]:
                df = filterDateFrame("本期收到的处置子公司的现金净额", path, conditions=("本期数",))
                if len(df) > 0:
                    addParagraph(document, "2、本期收到的处置子公司的现金净额", "paragraph")
                    dfToWord(document, df, style=2)

                addParagraph(document, "3、现金及现金等价物的构成", "paragraph")
                dfToWord(document, df1, style=2)
            else:
                addParagraph(document, "2、现金及现金等价物的构成", "paragraph")
                dfToWord(document, df1, style=2)
    else:
        addParagraph(document, "2、现金及现金等价物的构成", "paragraph")
        dfToWord(document, df1, style=2)


def addForeignCurrencyMonetaryItems(document, num, path):
    addTitle(document, "（{}）外币货币性项目".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "外币货币性项目", path, style=2,conditions=("期末外币余额",))


def addAssetsWithLimitedOwnershipOrRightOfUse(document, num, path):
    addTitle(document, "（{}）所有权和使用权受到限制的资产".format(to_chinese(num)), 2, True)
    excelTableToWord(document, "所有权或使用权受到限制的资产", path, style=2,conditions=("期末账面价值",))


# addTitle(document, "九、或有事项", 1, False)
# addTitle(document, "十、资产负债表日后事项", 1, False)
# addTitle(document, "十一、关联方关系及其交易", 1, False)
# addTitle(document, "（一）母公司基本情况", 2, True)
# addTitle(document, "（二）子公司情况", 2, True)
# addTitle(document, "（三）合营企业及联营企业情况", 2, True)
# addTitle(document, "（四）其他关联方", 2, True)
# addTitle(document, "（五）关联方交易", 2, True)


funcDict = {
    "货币资金": addMonetary,
    "交易性金融资产": addTradingFinancialAssets,
    "以公允价值计量且其变动计入当期损益的金融资产": addFinancialAssetsMeasuredAtFairValueWithChangesIncludedInCurrentProfitAndLoss,
    "衍生金融资产": addDerivativeFinancialAssets,
    "应收票据": addNotesReceivable,
    "应收账款": addAccountsReceivable,
    "应收款项融资": addReceivablesFinancing,
    "预付款项": addAdvancePayment,
    "其他应收款": addOtherReceivables,
    "存货": addInventory,
    "合同资产": addContractAssets,
    "持有待售资产": addAssetsHeldForSale,
    "一年内到期的非流动资产": addNonCurrentAssetsDueWithinOneYear,
    "其他流动资产": addOtherCurrentAssets,
    "债权投资": addDebtInvestment,
    "可供出售金融资产": addAvailableForSaleFinancialAssets,
    "其他债权投资": addOtherDebtInvestment,
    "持有至到期投资": addHeldToMaturityInvestment,
    "长期应收款": addlongTermReceivables,
    "长期股权投资": addLongTermEquityInvestment,
    "其他权益工具投资": addInvestmentInOtherEquityInstruments,
    "其他非流动金融资产": addOtherNonCurrentFinancialAssets,
    "投资性房地产": addInvestmentRealEstate,
    "固定资产": addFixedAssets,
    "在建工程": addConstructionInProgress,
    "生产性生物资产": addProductiveBiologicalAssets,
    "油气资产": addOilAndGasAssets,
    "使用权资产": addRightToUseAssets,
    "无形资产": addIntangibleAssets,
    "开发支出": addDevelopmentExpenditure,
    "商誉": addGoodwill,
    "长期待摊费用": addLongTermDeferredExpenses,
    "递延所得税资产": addDeferredTax,
    "其他非流动资产": addOtherNonCurrentAssets,
    "短期借款": addShortTermLoan,
    "交易性金融负债": addTradingFinancialLiabilities,
    "以公允价值计量且其变动计入当期损益的金融负债": addFinancialLiabilitiesMeasuredAtFairValueWithChangesIncludedInCurrentProfitAndLoss,
    "衍生金融负债": addDerivativeFinancialLiabilities,
    "应付票据": addNotesPayable,
    "应付账款": addAccountsPayable,
    "预收款项": addDepositReceived,
    "合同负债": addContractualLiabilities,
    "应付职工薪酬": addEmployeeCompensationPayable,
    "应交税费": addTaxPayable,
    "其他应付款": addOtherAccountsPayable,
    "持有待售负债": addLiabilitiesHeldForSale,
    "一年内到期的非流动负债": addNonCurrentLiabilitiesDueWithinOneYear,
    "其他流动负债": addOtherCurrentLiabilities,
    "长期借款": addLongTermLoan,
    "应付债券": addBondsPayable,
    "租赁负债": addLeaseLiabilities,
    "长期应付款": addLongTermAccountsPayable,
    "长期应付职工薪酬": addLongTermEmployeeCompensationPayable,
    "预计负债": addEstimatedLiabilities,
    "递延收益": addDeferredIncome,
    "递延所得税负债": addDeferredTax,
    "其他非流动负债": addOtherNonCurrentLiabilities,
    "实收资本（或股本）净额": addEquity,
    "股本": addEquity,
    "其他权益工具": addOtherEquityInstruments,
    "资本公积": addCapitalReserve,
    "其他综合收益": addOtherComprehensiveIncome,
    "专项储备": addSpecialReserve,
    "盈余公积": addSurplusReserve,
    "未分配利润": addUndistributedProfit,
    "其中：营业收入": addRevenueCost,
    "其中：营业成本": addRevenueCost,
    "一、营业收入": addRevenueCost,
    "减：营业成本": addRevenueCost,
    "税金及附加": addTaxesAndSurcharges,
    "销售费用": addsellingExpenses,
    "管理费用": addAdministrativeExpenses,
    "研发费用": addRDExpenses,
    "财务费用": addFinancialExpenses,
    "加：其他收益": addOtherIncome,
    "投资收益（损失以“-”号填列）": addincomeFromInvestment,
    "净敞口套期收益（损失以“-”号填列)": addNetExposureHedgingGains,
    "公允价值变动收益（损失以“-”号填列）": addIncomeFromChangesInFairValue,
    "信用减值损失（损失以“-”号填列）": addCreditImpairmentLoss,
    "资产减值损失（损失以“-”号填列）": addassetsImpairmentLoss,
    "资产处置收益（损失以“-”号填列）": addIncomeFromAssetDisposal,
    "加：营业外收入": addNonOperatingIncome,
    "减：营业外支出": addNonOperatingExpenses,
    "减：所得税费用": addIncomeTaxExpenses,
    "收到其他与经营活动有关的现金": addOtherCashReceivedRelatedToOperatingActivities,
    "支付其他与经营活动有关的现金": addOtherCashPaymentsRelatedToOperatingActivities,
    "收到其他与投资活动有关的现金": addOtherCashReceivedRelatedToInvestmentActivities,
    "支付其他与投资活动有关的现金": addOtherCashPaymentsRelatedToInvestmentActivities,
    "收到其他与筹资活动有关的现金": addOtherCashReceivedRelatedToFinancingActivities,
    "支付其他与筹资活动有关的现金": addOtherCashPaymentsRelatedToFinancingActivities,
}


# 添加报表附注
def addFsNote(document, path,context,isAll, assetsRecordsCombine, liabilitiesRecordsCombine, profitRecordsCombine,
              cashRecordsCombine,
              assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle, cashRecordsSingle):
    # 报告类型
    reportType = context["report_params"]["type"]


    combines = [assetsRecordsCombine, liabilitiesRecordsCombine, profitRecordsCombine, cashRecordsCombine]
    singles = [assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle, cashRecordsSingle]
    # 删除重复的递延所得税、营业收入

    if reportType == "合并":
        for combine in combines:
            for item in combine:
                if item["name"] == "递延所得税负债":
                    if not searchRecordItemByName("递延所得税资产", assetsRecordsCombine) is None:
                        continue
                elif item["name"] == "其中：营业成本":
                    if not searchRecordItemByName("其中：营业收入", profitRecordsCombine) is None:
                        continue
                if isAll:
                    if handleName(item["name"]) in funcDict:
                        funcDict[handleName(item["name"])](document, item["noteNum"], path, context)
                else:
                    if item["noteNum"] != "":
                        funcDict[handleName(item["name"])](document, item["noteNum"], path, context)
    else:
        for single in singles:
            for item in single:
                if item["name"] == "递延所得税负债":
                    if not searchRecordItemByName("递延所得税资产", assetsRecordsCombine) is None:
                        continue
                elif item["name"] == "其中：营业成本":
                    if not searchRecordItemByName("其中：营业收入", profitRecordsSingle) is None:
                        continue
                elif item["name"] == "减：营业成本":
                    if not searchRecordItemByName("一、营业收入", profitRecordsSingle) is None:
                        continue
                if isAll:
                    if handleName(item["name"]) in funcDict:
                        funcDict[handleName(item["name"])](document, item["noteNum"], path, context)
                else:
                    if item["noteNum"] != "":
                        funcDict[handleName(item["name"])](document, item["noteNum"], path, context)


# 获取单个报表最后一个编号
def getReportFormLastNum(reportForm):
    for i in range(len(reportForm) - 1, -1, -1):
        if reportForm[i]["noteNum"] != "":
            return reportForm[i]["noteNum"]
    return ""


# 获取所有财务报表的最后一个编号
def getLastNum(context, cashRecordsCombine, profitRecordsCombine, cashRecordsSingle, profitRecordsSingle):
    companyType = context["report_params"]["companyType"]
    # 报告类型
    reportType = context["report_params"]["type"]
    if companyType == "上市公司":
        if reportType == "合并":
            res = getReportFormLastNum(cashRecordsCombine)
            if res == "":
                return getReportFormLastNum(profitRecordsCombine)
        else:
            res = getReportFormLastNum(cashRecordsSingle)
            if res == "":
                return getReportFormLastNum(profitRecordsSingle)
    else:
        if reportType == "合并":
            return getReportFormLastNum(profitRecordsCombine)
        else:
            return getReportFormLastNum(profitRecordsSingle)

# # 或有事项
# def addContingencies(document, currentpath, context):
#     addParagraph(document, "(一) 未决诉讼", "paragraph")
#
#     addParagraph(document, "(二) 对外担保", "paragraph")
#     addParagraph(document, "1、为关联方提供的担保事项详见本财务报表附注十之说明。", "paragraph")
#     addParagraph(document, "2、公司及子公司为非关联方提供的担保事项", "paragraph")


def addOtherNoteAppended(document, currentpath,parentpath, context, assetsRecordsCombine,comparativeTable):
    companyType = context["report_params"]["companyType"]
    # 报告类型
    reportType = context["report_params"]["type"]

    if companyType == "国有企业":
        # 国有企业
        # 或有事项
        # 资产负债表日后事项
        # 关联方及其交易
        if reportType == "合并":
            addTitle(document, "九、或有事项", 1, False)
            addParagraph(document, "截止{}，本公司无需要披露的重大或有事项。".format(context["report_params"]["reportDate"]), "paragraph")
            addTitle(document, "十、资产负债表日后事项", 1, False)
            addEventsAfterBalanceSheetDate(document)
            addTitle(document, "十一、关联方及关联交易", 1, False)
            addRelationship(document, currentpath, context, assetsRecordsCombine)
            addTitle(document, "十二、母公司财务报表主要项目注释", 1, False)
            addParentCompanyNoteAppended(document, parentpath, context, comparativeTable)
        else:
            addTitle(document, "八、或有事项", 1, False)
            addParagraph(document, "截止{}，本公司无需要披露的重大或有事项。".format(context["report_params"]["reportDate"]), "paragraph")
            addTitle(document, "九、资产负债表日后事项", 1, False)
            addEventsAfterBalanceSheetDate(document)
            addTitle(document, "十、关联方及关联交易", 1, False)
            addRelationship(document, currentpath, context, assetsRecordsCombine)

    else:
        if reportType == "合并":
            addTitle(document, "七、合并范围变更", 1, False)
            addCombineRangeChange(document, currentpath)
            addTitle(document, "八、在其他主体中的权益", 1, False)
            addEquityInOtherEntities(document, currentpath, context)
            addTitle(document, "九、与金融工具相关的风险", 1, False)
            addRisksRelatedToFinancialInstruments(document, currentpath,context)
            addTitle(document, "十、公允价值的披露", 1, False)
            addFairValue(document, currentpath)
            addTitle(document, "十一、关联方及关联交易", 1, False)
            addRelationship(document, currentpath, context, assetsRecordsCombine)
            if context["noteAppend"]["shareBasedPayment"]:
                addTitle(document, "十二、股份支付", 1, False)
                addShareBasedPayment(document, currentpath, context)
                addTitle(document, "十三、承诺及或有事项", 1, False)
                addCommitmentContingency(document, currentpath, context)
                addTitle(document, "十四、资产负债表日后事项", 1, False)
                addEventsAfterBalanceSheetDate(document)
                addTitle(document, "十五、其他重要事项", 1, False)
                addOtherImportantEvent(document, currentpath, context)
                addTitle(document, "十六、母公司财务报表主要项目注释", 1, False)
                addParentCompanyNoteAppended(document, parentpath, context, comparativeTable)
                addTitle(document, "十七、其他补充资料", 1, False)
                addOtherSupplementaryInformation(document, currentpath)
            else:
                addTitle(document, "十二、承诺及或有事项", 1, False)
                addCommitmentContingency(document, currentpath, context)
                addTitle(document, "十三、资产负债表日后事项", 1, False)
                addEventsAfterBalanceSheetDate(document)
                addTitle(document, "十四、其他重要事项", 1, False)
                addOtherImportantEvent(document, currentpath, context)
                addTitle(document, "十五、母公司财务报表主要项目注释", 1, False)
                addParentCompanyNoteAppended(document, parentpath, context, comparativeTable)
                addTitle(document, "十六、其他补充资料", 1, False)
                addOtherSupplementaryInformation(document, currentpath)

        else:
            addTitle(document, "七、与金融工具相关的风险", 1, False)
            addRisksRelatedToFinancialInstruments(document, currentpath,context)
            addTitle(document, "八、公允价值的披露", 1, False)
            addFairValue(document, currentpath)
            addTitle(document, "九、关联方及关联交易", 1, False)
            addRelationship(document, currentpath, context, assetsRecordsCombine)
            if context["noteAppend"]["shareBasedPayment"]:
                addTitle(document, "十、股份支付", 1, False)
                addShareBasedPayment(document, currentpath, context)
                addTitle(document, "十一、承诺及或有事项", 1, False)
                addCommitmentContingency(document, currentpath, context)
                addTitle(document, "十二、资产负债表日后事项", 1, False)
                addEventsAfterBalanceSheetDate(document)
                addTitle(document, "十三、其他重要事项", 1, False)
                addOtherImportantEvent(document, currentpath, context)
                addTitle(document, "十四、其他补充资料", 1, False)
                addOtherSupplementaryInformation(document, currentpath)

            else:
                addTitle(document, "十、承诺及或有事项", 1, False)
                addCommitmentContingency(document, currentpath, context)
                addTitle(document, "十一、资产负债表日后事项", 1, False)
                addEventsAfterBalanceSheetDate(document)
                addTitle(document, "十二、其他重要事项", 1, False)
                addOtherImportantEvent(document, currentpath, context)
                addTitle(document, "十三、其他补充资料", 1, False)
                addOtherSupplementaryInformation(document, currentpath)


# 其他补充资料
def addOtherSupplementaryInformation(document, path):
    addParagraph(document, "(一) 非经常性损益", "paragraph")
    excelTableToWord(document, "非经常性损益上市公司", path, style=2,conditions=("金额",))

    addParagraph(document, "(二) 净资产收益率及每股收益", "paragraph")
    df = pd.read_excel(path,sheet_name="净资产收益率及每股收益上市公司")
    dfToWord(document,df,style=2)


# 添加母公司财务报表注释
def addParentCompanyNoteAppended(document,parentPath,context,comparativeTable):
    # 公司类型
    companyType = context["report_params"]["companyType"]
    assetsRecordsSingle = searchModel(companyType, "单体", "资产表", comparativeTable)
    liabilitiesRecordsSingle = searchModel(companyType, "单体", "负债表", comparativeTable)
    profitRecordsSingle = searchModel(companyType, "单体", "利润表", comparativeTable)

    singles = [assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle]
    for single in singles:
        for item in single:
            if item["name"] == "其中：营业成本":
                if not searchRecordItemByName("其中：营业收入", profitRecordsSingle) is None:
                    continue
            elif item["name"] == "减：营业成本":
                if not searchRecordItemByName("一、营业收入", profitRecordsSingle) is None:
                    continue
            else:
                if item["noteNum"] != "":
                    funcDict[handleName(item["name"])](document, item["noteNum"], parentPath, context)


def addSegmentInformation(document, path):
    addParagraph(document, "1、报告分部的确定依据与会计政策", "paragraph")
    addParagraph(document,
                 "公司以内部组织结构、管理要求、内部报告制度等为依据确定报告分部，并以行业分部/产品分部为基础确定报告分部。分别对   业务、   业务及   业务等的经营业绩进行考核。与各分部共同使用的资产、负债按照规模比例在不同的分部之间分配。",
                 "paragraph")
    addParagraph(document, "本公司以地区分部为基础确定报告分部，主营业务收入、主营业务成本按最终实现销售地进行划分，资产和负债按经营实体所在地进行划分。", "paragraph")

    addParagraph(document, "2、报告分部的财务信息", "paragraph")
    df = pd.read_excel(path,sheet_name="分部信息业务分部期末数")
    dfToWord(document,df,style=2)
    addParagraph(document, "（续）", "paragraph")
    df = pd.read_excel(path, sheet_name="分部信息业务分部期初数")
    dfToWord(document, df, style=2)


    addParagraph(document, "3、按收入来源地划分的对外交易收入", "paragraph")
    excelTableToWord(document, "按收入来源地划分的对外交易收入", path, style=2,conditions=("本期数","上期数"))

    addParagraph(document, "4、按资产所在地划分的非流动资产", "paragraph")
    excelTableToWord(document, "按资产所在地划分的非流动资产", path, style=2,conditions=("本期数","上期数"))

    addParagraph(document, "5、本公司对主要客户的依赖程度", "paragraph")
    excelTableToWord(document, "本公司对主要客户的依赖程度", path, style=2,conditions=("销售额",))


# 其他重要事项
def addOtherImportantEvent(document, path, context):
    if context["noteAppend"]["SegmentInformation"]:
        addTitle(document, "(一）分部信息", 2, True)
        addSegmentInformation(document, path)


# 资产负债表日后事项
def addEventsAfterBalanceSheetDate(document):
    addParagraph(document, "本公司无需要披露的资产负债表日后事项。", "paragraph")


# 承诺及或有事项
def addCommitmentContingency(document, path, context):
    addTitle(document, "(一）重大承诺事项", 2, True)
    addParagraph(document, "1、资本承诺", "paragraph")
    excelTableToWord(document, "资本承诺上市公司", path, style=2)
    addParagraph(document, "截至{}，本公司无需要披露的重大承诺事项。".format(context["report_params"]["reportDate"]), "paragraph")
    addTitle(document, "(二）或有事项", 2, True)
    addParagraph(document, "截至{}，本公司无需要披露的重大或有事项。".format(context["report_params"]["reportDate"]), "paragraph")


# 公允价值的披露
def addFairValue(document, path):
    addTitle(document, "(一）以公允价值计量的资产和负债的年末公允价值", 2, True)
    addParagraph(document,
                 "下表列示了本公司在每个资产负债表日持续和非持续以公允价值计量的资产和负债于本报告期末的公允价值信息及其公允价值计量的层次。公允价值计量结果所属层次取决于对公允价值计量整体而言具有重要意义的最低层次的输入值。三个层次输入值的定义如下：",
                 "paragraph")
    addParagraph(document, "第一层次：相同资产或负债在活跃市场上未经调整的报价。", "paragraph")
    addParagraph(document, "第二层次：除第一层次输入值外相关资产或负债直接或间接可观察的输入值。", "paragraph")
    addParagraph(document, "第三层次：相关资产或负债的不可观察输入值。", "paragraph")
    excelTableToWord(document, "以公允价值计量的资产和负债的期末公允价值明细情况上市公司", path, style=2,conditions=("合计",))

    addTitle(document, "(二）持续和非持续第一层次公允价值计量项目市价的确定依据", 2, True)
    addParagraph(document, "本公司以活跃市场报价作为第一层次金融资产的公允价值。", "paragraph")

    addTitle(document, "(三）持续和非持续第二层次公允价值计量项目，采用的估值技术和重要参数的定性及定量信息", 2, True)
    addParagraph(document, "衍生金融资产及衍生金融负债为本集团与金融机构签订的多个 IRS、CCS，本集团使用金融机构提供的报价作为估值依据。", "paragraph")

    addTitle(document, "(四）持续和非持续第三层次公允价值计量项目，采用的估值技术和重要参数的定性及定量信息", 2, True)
    addParagraph(document,
                 "持续第三层次公允价值计量的其他非流动金融资产主要为本公司持有的未上市股权投资。本公司采用估值技术进行了公允价值计量，主要采用了上市公司比较法的估值技术，参考类似证券的股票价格并考虑流动性折扣。",
                 "paragraph")
    addParagraph(document, "持续第三层次公允价值计量的交易性金融资产主要为本公司持有的理财产品。本公司采用约定的预期收益率计算的未来现金流量折现的方法估算公允价值。", "paragraph")
    addParagraph(document,
                 "非持续第三层次公允价值计量的持有待售资产主要为本集团持有的拟出售项目。本集团按账面价值与公允价值减去出售费用后净额之孰低者对持有待售资产进行初始计量和后续计量。本公司主要参考独立合格专业评估师的评估报告，该评估主要采用了现金流量折现法及市场比较法的估值技术。",
                 "paragraph")

    addTitle(document, "(五）持续的第三层次公允价值计量项目，期初与期末账面价值间的调节信息及不可观察参数敏感性分析", 2, True)
    excelTableToWord(document, "第三层次公允价值计量项目期初与期末账面价值间的调节信息上市公司", path, style=2)

    addTitle(document, "(六）持续的公允价值计量项目，本期内发生各层级之间转换的，转换的原因及确定转换时点的政策", 2, True)
    addParagraph(document, "本公司以导致各层次之间转换的事项发生日为确认各层次之间转换的时点。本公司的金融资产及金融负债的公允价值计量未发生各层级之间的转换。", "paragraph")

    addTitle(document, "(七）本期发生的估值技术变更及变更原因", 2, True)
    addParagraph(document, "本公司金融工具的公允价值计量方法并未发生改变。", "paragraph")

    addTitle(document, "(八）不以公允价值计量的金融资产和金融负债的公允价值情况", 2, True)
    addParagraph(document, "本公司以摊余成本计量的金融资产和金融负债主要包括：应收款项、短期借款、应付款项、长期借款和应付债券等。", "paragraph")
    addParagraph(document, "除下述金融资产和金融负债以外，其他不以公允价值计量的金融资产和金融负债的账面价值与公允价值差异很小。", "paragraph")
    excelTableToWord(document, "不以公允价值计量的金融资产和金融负债的公允价值情况上市公司", path, style=2,conditions=("期末账面价值","期初账面价值"))


# 金融工具相关的风险
def addRisksRelatedToFinancialInstruments(document, path,context):
    addParagraph(document,
                 "本公司从事风险管理的目标是在风险和收益之间取得平衡，将风险对本公司经营业绩的负面影响降至最低水平，使股东和其他权益投资者的利益最大化。基于该风险管理目标，本公司风险管理的基本策略是确认和分析本公司面临的各种风险，建立适当的风险承受底线和进行风险管理，并及时可靠地对各种风险进行监督，将风险控制在限定的范围内。",
                 "paragraph")
    addParagraph(document, "本公司在日常活动中面临各种与金融工具相关的风险，主要包括信用风险、流动性风险及市场风险。管理层已审议并批准管理这些风险的政策，概括如下。", "paragraph")
    addTitle(document, "(一）信用风险", 2, True)
    addParagraph(document, "信用风险，是指金融工具的一方不能履行义务，造成另一方发生财务损失的风险。", "paragraph")
    addParagraph(document, "1、信用风险管理实务", "paragraph")
    addParagraph(document, "(1)信用风险的评价方法", "paragraph")
    addParagraph(document,
                 "公司在每个资产负债表日评估相关金融工具的信用风险自初始确认后是否已显著增加。在确定信用风险自初始确认后是否显著增加时，公司考虑在无须付出不必要的额外成本或努力即可获得合理且有依据的信息，包括基于历史数据的定性和定量分析、外部信用风险评级以及前瞻性信息。公司以单项金融工具或者具有相似信用风险特征的金融工具组合为基础，通过比较金融工具在资产负债表日发生违约的风险与在初始确认日发生违约的风险，以确定金融工具预计存续期内发生违约风险的变化情况。",
                 "paragraph")
    addParagraph(document, "当触发以下一个或多个定量、定性标准时，公司认为金融工具的信用风险已发生显著增加：", "paragraph")
    addParagraph(document, "1) 定量标准主要为资产负债表日剩余存续期违约概率较初始确认时上升超过一定比例；", "paragraph")
    addParagraph(document, "2) 定性标准主要为债务人经营或财务情况出现重大不利变化、现存的或预期的技术、市场、经济或法律环境变化并将对债务人对公司的还款能力产生重大不利影响等；", "paragraph")
    addParagraph(document, "3) 上限标准为债务人合同付款(包括本金和利息)逾期超过90天。", "paragraph")
    addParagraph(document, "(2) 违约和已发生信用减值资产的定义", "paragraph")
    addParagraph(document, "当金融工具符合以下一项或多项条件时，公司将该金融资产界定为已发生违约，其标准与已发生信用减值的定义一致：", "paragraph")
    addParagraph(document, "1)定量标准", "paragraph")
    addParagraph(document, "债务人在合同付款日后逾期超过90天仍未付款；", "paragraph")
    addParagraph(document, "2) 定性标准", "paragraph")
    addParagraph(document, "①债务人发生重大财务困难；", "paragraph")
    addParagraph(document, "②债务人违反合同中对债务人的约束条款；", "paragraph")
    addParagraph(document, "③债务人很可能破产或进行其他财务重组；", "paragraph")
    addParagraph(document, "④ 债权人出于与债务人财务困难有关的经济或合同考虑，给予债务人在任何其他情况下都不会做出的让步。", "paragraph")
    addParagraph(document, "2、预期信用损失的计量", "paragraph")
    addParagraph(document,
                 "预期信用损失计量的关键参数包括违约概率、违约损失率和违约风险敞口。公司考虑历史统计数据(如交易对手评级、担保方式及抵质押物类别、还款方式等)的定量分析及前瞻性信息，建立违约概率、违约损失率及违约风险敞口模型。",
                 "paragraph")
    addParagraph(document, "3、信用风险敞口及信用风险集中度", "paragraph")
    addParagraph(document,
                 "信用风险主要产生于银行存款、应收票据、应收账款、应收款项融资、其他应收款、合同资产、债权投资、长期应收款以及财务担保合同等以及未纳入减值评估范围的以公允价值计量且其变动计入当期损益的债务工具投资和衍生金融资产等。于资产负债表日，本公司金融资产的账面价值已代表其最大信用风险敞口。",
                 "paragraph")
    addParagraph(document, "本公司银行存款主要存放于国有银行和其他大中型上市银行以及拥有较高信用评级的国内和海外银行，本公司认为其不存在重大的信用风险，不会产生因对方单位违约而导致的任何重大损失。",
                 "paragraph")
    addParagraph(document,
                 "此外，对于应收票据、应收账款、应收款项融资、其他应收款、合同资产和长期应收款，本公司设定相关政策以控制信用风险敞口。本公司基于对客户的财务状况、信用记录及其他因素诸如目前市场状况等评估客户的信用资质并设置相应信用期。本公司会定期对客户信用记录进行监控，对于信用记录不良的客户，本集团会采用书面催款、缩短信用期或取消信用期等方式，以确保本公司的整体信用风险在可控的范围内。",
                 "paragraph")
    addParagraph(document,
                 "（提示：适用信用风险集中）由于本公司仅与经认可的且信用良好的第三方进行交易，所以无需担保物。信用风险集中按照客户进行管理。截至2019年12月31日，本公司存在一定的信用集中风险，本公司应收账款的  %(2018年12月31日：  %)源于余额前五名客户。本公司对应收账款余额未持有任何担保物或其他信用增级。",
                 "paragraph")
    addParagraph(document,
                 "（提示：适用不存在信用集中风险）由于本公司的应收账款风险点分布于多个合作方和多个客户，截至2019年12月31日，本公司应收账款的  %(2018年12月31日：  %)源于余额前五名客户，本公司不存在重大的信用集中风险。",
                 "paragraph")
    addParagraph(document, "本公司所承受的最大信用风险敞口为资产负债表中每项金融资产的账面价值。", "paragraph")

    addTitle(document, "(二）流动性风险", 2, True)
    addParagraph(document,
                 "流动性风险，是指本公司在履行以交付现金或其他金融资产的方式结算的义务时发生资金短缺的风险。流动性风险可能源于无法尽快以公允价值售出金融资产；或者源于对方无法偿还其合同债务；或者源于提前到期的债务；或者源于无法产生预期的现金流量。",
                 "paragraph")
    addParagraph(document,
                 "为控制该项风险，本公司综合运用票据结算、银行借款等多种融资手段，并采取长、短期融资方式适当结合，优化融资结构的方法，保持融资持续性与灵活性之间的平衡。本公司已从多家商业银行取得银行授信额度以满足营运资金需求和资本开支。",
                 "paragraph")
    addParagraph(document, "1、期末金融负债按剩余到期日分类", "paragraph")
    excelTableToWord(document, "金融负债按剩余到期日分类期末数上市公司", path, style=2,conditions=("账面价值",))
    addParagraph(document, "2、期初金融负债按剩余到期日分类", "paragraph")
    excelTableToWord(document, "金融负债按剩余到期日分类期初数上市公司", path, style=2,conditions=("账面价值",))

    addTitle(document, "(三）市场风险", 2, True)
    addParagraph(document, "市场风险，是指金融工具的公允价值或未来现金流量因市场价格变动而发生波动的风险。市场风险主要包括利率风险和外汇风险。", "paragraph")
    addParagraph(document, "1、利率风险", "paragraph")
    addParagraph(document,
                 "利率风险，是指金融工具的公允价值或未来现金流量因市场利率变动而发生波动的风险。固定利率的带息金融工具使本公司面临公允价值利率风险，浮动利率的带息金融工具使本公司面临现金流量利率风险。本公司因利率变动引起金融工具公允价值变动的风险主要与固定利率的应付债券、固定利率的银行借款、固定利率的债权投资、一年内到期的非流动资产、其他流动资产、衍生金融资产/负债-利率掉期合同有关，而现金流量变动的风险主要与浮动利率银行借款有关。本公司根据当时的市场环境来决定固定利率及浮动利率合同的相对比例。于{}，本公司长短期带息债务主要为人民币计价的浮动利率合同，金额为人民币***元。交易性金融资产及其他非流动金融资产中的非上市信托产品投资主要为人民币计价的浮动收益合同，成本金额为人民币***元。".format(
                     context["report_params"]["reportDate"]), "paragraph")
    addParagraph(document,
                 "本公司密切关注利率变动对本公司利率风险的影响。本公司目前并未采取利率对冲政策。但管理层负责监控利率风险，并将于需要时考虑对冲重大利率风险。利率上升会增加新增带息债务的成本以及本集团尚未付清的以浮动利率计息的带息债务的利息费用，并对本公司的财务业绩产生重大的不利影响，管理层会依据最新的市场状况及时做出调整，这些调整可能是进行利率互换的安排来降低利率风险。","paragraph")
    addParagraph(document, "利率敏感性分析如下：", "paragraph")
    excelTableToWord(document, "利率敏感性分析上市公司", path, style=2,conditions=("对本期净利润的影响",))
    addParagraph(document, "2、外汇风险", "paragraph")
    addParagraph(document,
                 "外汇风险，是指金融工具的公允价值或未来现金流量因外汇汇率变动而发生波动的风险。本公司面临的汇率变动的风险主要与本公司外币货币性资产和负债有关。对于外币资产和负债，如果出现短期的失衡情况，本公司会在必要时按市场汇率买卖外币，以确保将净风险敞口维持在可接受的水平。/本公司于中国内地经营，且主要活动以人民币计价。因此，本公司所承担的外汇变动市场风险不重大。",
                 "paragraph")
    addParagraph(document, "本公司持有的外币金融资产和外币金融负债折算成人民币的金额列示如下：", "paragraph")
    excelTableToWord(document, "外币货币性项目上市公司", path, style=2)
    addParagraph(document, "外汇风险敏感性分析", "paragraph")
    excelTableToWord(document, "外汇风险敏感性分析", path, style=2,conditions=("对本期净利润的影响",))


# 合并范围变更
def addCombineRangeChange(document, path):
    addTitle(document, "(一）非同一控制下企业合并", 2, True)
    df = pd.read_excel(path, sheet_name="非同一控制下企业合并上市公司")

    if len(df)>0:
        addTitle(document, "1、本期发生的非同一控制下企业合并", 2, True)
        dfToWord(document,df,style=2)
        addTitle(document, "2、合并成本及商誉", 2, True)
        df = pd.read_excel(path, sheet_name="合并成本及商誉")
        dfToWord(document, df, style=2)
        addTitle(document, "3、被购买方于购买日可辨认资产、负债", 2, True)
        df = pd.read_excel(path, sheet_name="被购买方于购买日可辨认资产负债")
        dfToWord(document, df, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")


    addTitle(document, "(二）同一控制下企业合并", 2, True)
    df = pd.read_excel(path, sheet_name="同一控制下企业合并")
    if len(df)>0:
        addTitle(document, "1、本期发生的同一控制下企业合并", 2, True)
        dfToWord(document,df,style=2)
        addTitle(document, "2、合并成本", 2, True)
        df = pd.read_excel(path, sheet_name="合并成本")
        dfToWord(document, df, style=2)
        addTitle(document, "3、合并日被合并方资产、负债的账面价值", 2, True)
        df = pd.read_excel(path, sheet_name="合并日被合并方资产负债的账面价值")
        dfToWord(document, df, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")

    addTitle(document, "(三）处置子公司", 2, True)
    df = pd.read_excel(path,sheet_name="单次处置对子公司投资即丧失控制权")
    if len(df)>0:
        addTitle(document, "1、单次处置对子公司投资即丧失控制权", 2, True)
        dfToWord(document,df,style=2)
        df = pd.read_excel(path,sheet_name="多次处置构成一揽子交易")
        if len(df)>0:
            addTitle(document, "2、通过多次交易分步处置对子公司投资且在本期丧失控制权", 2, True)
            addTitle(document, "（1）构成一揽子交易", 2, True)
            dfToWord(document, df, style=2)
            df = pd.read_excel(path,sheet_name="多次处置不构成一揽子交易")
            if len(df)>0:
                addTitle(document, "（2）不构成一揽子交易", 2, True)
                dfToWord(document, df, style=2)
        else:
            df = pd.read_excel(path, sheet_name="多次处置不构成一揽子交易")
            if len(df) > 0:
                addTitle(document, "2、通过多次交易分步处置对子公司投资且在本期丧失控制权", 2, True)
                addTitle(document, "（1）不构成一揽子交易", 2, True)
                dfToWord(document, df, style=2)
    else:
        df = pd.read_excel(path, sheet_name="多次处置构成一揽子交易")
        if len(df)>0:
            addTitle(document, "1、通过多次交易分步处置对子公司投资且在本期丧失控制权", 2, True)
            addTitle(document, "（1）构成一揽子交易", 2, True)
            dfToWord(document, df, style=2)
            df = pd.read_excel(path,sheet_name="多次处置不构成一揽子交易")
            if len(df)>0:
                addTitle(document, "（2）不构成一揽子交易", 2, True)
                dfToWord(document, df, style=2)
        else:
            df = pd.read_excel(path, sheet_name="多次处置不构成一揽子交易")
            if len(df) > 0:
                addTitle(document, "1、通过多次交易分步处置对子公司投资且在本期丧失控制权", 2, True)
                addTitle(document, "（1）不构成一揽子交易", 2, True)
                dfToWord(document, df, style=2)
            else:
                addParagraph(document, "不适用", "paragraph")



    addTitle(document, "(四）其他原因的合并范围变动", 2, True)
    df1 = filterDateFrame("其他合并范围增加",path,conditions=("出资额",))
    df2 = filterDateFrame("其他合并范围减少",path,conditions=("处置日净资产",))
    if len(df1)==0 and len(df2)==0:
        addParagraph(document, "不适用", "paragraph")
    else:
        addTitle(document, "1、合并范围增加", 2, True)
        dfToWord(document,df1,style=2)
        addTitle(document, "2、合并范围减少", 2, True)
        dfToWord(document, df2, style=2)


# 在其他主体中的权益
def addEquityInOtherEntities(document, path, context):
    # 报告日期
    reportDate = context["report_params"]["reportDate"]
    # 获取报告起始日
    startYear = reportDate[:4]

    addTitle(document, "（一）在子公司中的权益", 2, True)

    addParagraph(document, "1、企业集团的构成", "paragraph")
    df = pd.read_excel(path, sheet_name="企业集团的构成上市公司")
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["子公司名称", "主要经营地", "注册地", "业务性质", "持股比例（%）", "nan", "取得方式"],
                  ["nan", "nan", "nan", "nan", "直接", "间接", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 6]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
        addParagraph(document, "①在子公司的持股比例不同于表决权比例的说明", "paragraph")
        addParagraph(document, "②持有半数或以下表决权但仍控制被投资单位，以及持有半数以上表决权但不控制被投资单位的依据", "paragraph")
        addParagraph(document, "③对于纳入合并范围的重要的结构化主体，控制的依据", "paragraph")
        addParagraph(document, "④确定公司是代理人还是委托人的依据", "paragraph")

    addParagraph(document, "2、重要的非全资子公司", "paragraph")
    excelTableToWord(document, "重要的非全资子公司上市公司", path, style=2,conditions=("期末少数股东权益余额",))

    addParagraph(document, "3、重要的非全资子公司的主要财务信息", "paragraph")
    # 资产负债信息
    addParagraph(document, "（1）资产和负债情况", "paragraph")
    df = filterDateFrame("重要非全资子企业期末资产负债上市公司",path,conditions=("期末资产合计",))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["单位名称", "期末数", "nan", "nan", "nan", "nan", "nan"],
                  ["nan", "流动资产", "非流动资产", "资产合计", "流动负债", "非流动负债", "负债合计"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=3)
        addParagraph(document, "(续上表)：".format(startYear), "paragraph")
    df = filterDateFrame("重要非全资子企业期初资产负债上市公司", path, conditions=("期末资产合计",))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["单位名称", "期初数", "nan", "nan", "nan", "nan", "nan"],
                  ["nan", "流动资产", "非流动资产", "资产合计", "流动负债", "非流动负债", "负债合计"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=3)

    # 损益及现金流信息
    addParagraph(document, "（2）损益和现金流量情况", "paragraph")
    df = filterDateFrame("重要非全资子企业本期损益和现金流量情况上市公司",path,conditions=("本期净利润",))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["单位名称", "本期数", "nan", "nan", "nan"], ["nan", "营业收入", "净利润", "综合收益总额", "经营活动现金流量"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=3)
        addParagraph(document, "(续上表)：".format(startYear), "paragraph")
    df = filterDateFrame("重要非全资子企业上期损益和现金流量情况上市公司", path, conditions=("本期净利润",))
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["单位名称", "上期数", "nan", "nan", "nan"], ["nan", "营业收入", "净利润", "综合收益总额", "经营活动现金流量"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTableTitle(table, titles)
        addContentToCombineTitle(document, dc, table, titleLength, style=3)

    addParagraph(document, "4、使用企业集团资产和清偿企业集团债务的重大限制", "paragraph")
    addParagraph(document, "5、向纳入合并财务报表范围的结构化主体提供的财务支持或其他支持", "paragraph")

    addTitle(document, "（二）在子公司的所有者权益份额发生变化但仍控制子公司的交易", 2, True)
    addParagraph(document, "1、在子公司的所有者权益份额发生变化的情况说明", "paragraph")
    excelTableToWord(document, "在子公司的所有者权益份额发生变化的情况说明上市公司", path, style=2,conditions=("变动前持股比例","变动后持股比例"))
    addParagraph(document, "2、交易对于少数股东权益及归属于母公司所有者权益的影响", "paragraph")
    df = pd.read_excel(path,sheet_name="交易对于少数股东权益及归属于母公司所有者权益的影响上市公司")
    dfToWord(document,df,style=2)


    addTitle(document, "（三）在合营企业或联营企业中的权益", 2, True)
    addParagraph(document, "1、重要的合营企业或联营企业", "paragraph")
    df = pd.read_excel(path, sheet_name="重要的合营企业或联营企业上市公司")
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["合营企业或联营企业名称", "主要经营地", "注册地", "业务性质", "持股比例（%）", "nan", "对合营企业或联营企业投资的会计处理方法"],
                  ["nan", "nan", "nan", "nan", "直接", "间接", "nan"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [[1, 6]])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)

    addParagraph(document, "2、重要合营企业的主要财务信息", "paragraph")

    df = pd.read_excel(path,sheet_name="重要合营企业财务信息本期数上市公司")
    if len(df)>0:
        addParagraph(document, "本期数：", "paragraph")
        dfToWord(document,df,style=2)
    else:
        addParagraph(document, "不适用", "paragraph")

    df =  pd.read_excel(path,sheet_name="重要合营企业财务信息上期数上市公司")
    if len(df)>0:
        addParagraph(document, "上期数：", "paragraph")
        dfToWord(document,df,style=2)

    addParagraph(document, "3、重要联营企业的主要财务信息", "paragraph")
    df = pd.read_excel(path, sheet_name="重要联营企业财务信息本期数上市公司")
    if len(df)>0:
        addParagraph(document, "本期数：", "paragraph")
        dfToWord(document,df,style=2)
    else:
        addParagraph(document, "不适用", "paragraph")

    df =  pd.read_excel(path,sheet_name="重要联营企业财务信息上期数上市公司")
    if len(df)>0:
        addParagraph(document, "上期数：", "paragraph")
        dfToWord(document,df,style=2)


    addParagraph(document, "4、不重要的合营企业和联营企业的汇总财务信息", "paragraph")
    df = pd.read_excel(path,sheet_name="不重要合营企业和联营企业的汇总信息上市公司")
    dfToWord(document,df,style=2)

    addParagraph(document, "5、合营企业或联营企业发生的超额亏损", "paragraph")
    df = pd.read_excel(path, sheet_name="合营企业或联营企业发生的超额亏损上市公司")
    dfToWord(document, df, style=2)

    addTitle(document, "（四）重要的共同经营", 2, True)
    df = pd.read_excel(path, sheet_name="重要的共同经营上市公司")
    if len(df)>0:
        dc = df.to_dict("split")
        titles = [["共同经营名称", "主要经营地", "注册地", "业务性质", "持股比例（%）", "nan"],
                  ["nan", "nan", "nan", "nan", "直接", "间接"]]
        titleLength = len(titles)
        rowLength = len(dc["index"]) + titleLength
        columnLength = len(dc["columns"])
        table = createBorderedTable(document, rowLength, columnLength)
        addCombineTitleSpecialReceivable(titles, table, [])
        addContentToCombineTitle(document, dc, table, titleLength, style=2)
    else:
        addParagraph(document, "不适用", "paragraph")

    addTitle(document, "（五）在未纳入合并财务报表范围的结构化主体中的权益", 2, True)
    addParagraph(document, "不适用", "paragraph")


# 添加关联方及其交易
def addRelationship(document, path, context, assetsRecordsCombine):
    # 添加关联方
    addRelatedParties(document, path, context, assetsRecordsCombine)
    #     添加关联方交易
    addRelatedPartyTransactions(document, path, context)
    #     添加关联方余额
    addBalanceOfRelatedParties(document, path)
    # 添加关联方承诺
    addCommitmentOfRelatedParties(document, path, context)


# 添加关联方情况
def addRelatedParties(document, path, context, assetsRecordsCombine):
    companyType = context["report_params"]["companyType"]
    # 报告类型
    reportType = context["report_params"]["type"]

    addTitle(document, "(一）关联方情况", 2, True)
    addTitle(document, "1、母公司情况", 2, True)
    df = pd.read_excel(path,sheet_name="母公司基本情况")
    if len(df)>0:
        dfToWord(document,df,style=2)
    else:
        addParagraph(document, "不适用", "paragraph")
    if reportType == "合并":
        item = searchRecordItemByName("长期股权投资", assetsRecordsCombine)
        if companyType == "上市公司":
            addTitle(document, "2、子公司情况", 2, True)
            addParagraph(document, "详见附注八、1、在子公司中的权益。", "paragraph")
            if item["noteNum"] != "":
                addTitle(document, "3、合营企业和联营企业情况", 2, True)
                addParagraph(document, "本公司重要的合营和联营企业详见附注八、3、在合营企业或联营企业中的权益。", "paragraph")
                df = pd.read_excel(path, sheet_name="其他关联方情况")
                if len(df) > 0:
                    addTitle(document, "4、其他关联方情况", 2, True)
                    dfToWord(document, df, style=2)

            else:
                df = pd.read_excel(path, sheet_name="其他关联方情况")
                if len(df) > 0:
                    addTitle(document, "3、其他关联方情况", 2, True)
                    dfToWord(document, df, style=2)

        else:
            addTitle(document, "2、子公司情况", 2, True)
            addParagraph(document, "详见附注七、企业合并及合并财务报表。", "paragraph")
            if item["noteNum"] != "":
                addTitle(document, "3、合营企业和联营企业情况", 2, True)
                addParagraph(document, "详见附注八、（{}）、长期股权投资。".format(to_chinese(item["noteNum"])), "paragraph")
                df = pd.read_excel(path, sheet_name="其他关联方情况")
                if len(df) > 0:
                    addTitle(document, "4、其他关联方情况", 2, True)
                    dfToWord(document, df, style=2)

            else:
                df = pd.read_excel(path, sheet_name="其他关联方情况")
                if len(df) > 0:
                    addTitle(document, "3、其他关联方情况", 2, True)
                    dfToWord(document, df, style=2)
    else:
        df = pd.read_excel(path, sheet_name="其他关联方情况")
        if len(df) > 0:
            addTitle(document, "2、其他关联方情况", 2, True)
            dfToWord(document, df, style=2)


# 添加关联方交易
def addRelatedPartyTransactions(document, path, context):
    companyType = context["report_params"]["companyType"]

    addTitle(document, "(二）关联方交易", 2, True)
    addTitle(document, "1、定价政策", 2, True)
    addParagraph(document, context["noteAppend"]["relationTransactionPrice"], "paragraph")

    addTitle(document, "2、购销商品、提供和接受劳务的关联交易", 2, True)
    df1 = filterDateFrame("采购商品接收劳务",path,conditions=("本期数","上年同期数"))
    df2 = filterDateFrame("出售商品提供劳务",path,conditions=("本期数","上年同期数"))
    if len(df1)==0 and len(df2)==0:
        addParagraph(document, "不适用", "paragraph")
    else:
        if  len(df1)>0:
            addParagraph(document, "采购商品/接受劳务情况", "paragraph")
            dfToWord(document,df1,style=2)
        if len(df2)>0:
            addParagraph(document, "出售商品/提供劳务情况", "paragraph")
            dfToWord(document, df2, style=2)

    addTitle(document, "3、关联租赁情况", 2, True)
    df1 = filterDateFrame("本公司作为出租方", path, conditions=("本期数", "上年同期数"))
    df2 = filterDateFrame("本公司作为承租方", path, conditions=("本期数", "上年同期数"))
    if len(df1) == 0 and len(df2) == 0:
        addParagraph(document, "不适用", "paragraph")
    else:
        if len(df1) > 0:
            addParagraph(document, "本公司作为出租方", "paragraph")
            dfToWord(document, df1, style=2)
        if len(df2) > 0:
            addParagraph(document, "本公司作为承租方", "paragraph")
            dfToWord(document, df2, style=2)
            df3 = filterDateFrame("本公司作为承租方当期承担的租赁负债利息支出", path, conditions=("本期数", "上年同期数"))
            if len(df3)>0:
                addParagraph(document, "本公司作为承租方当期承担的租赁负债利息支出", "paragraph")
                dfToWord(document,df3,style=2)


    addTitle(document, "4、关联担保情况", 2, True)
    df1 = filterDateFrame("本公司作为担保方", path, conditions=("担保金额",))
    df2 = filterDateFrame("本公司作为被担保方", path, conditions=("担保金额", ))
    if len(df1)==0 and len(df2)==0:
        addParagraph(document, "不适用", "paragraph")
    else:
        if  len(df1)>0:
            addParagraph(document, "本公司作为担保方", "paragraph")
            dfToWord(document,df1,style=2)
        if len(df2)>0:
            addParagraph(document, "本公司作为被担保方", "paragraph")
            dfToWord(document, df2, style=2)

    addTitle(document, "5、关联方资金拆借", 2, True)
    df = filterDateFrame("关联方资金拆借", path, conditions=("拆借金额", ))
    dfToWord(document,df,style=2)


    if companyType == "上市公司":
        addTitle(document, "6、关键管理人员薪酬", 2, True)
        excelTableToWord(document, "关键管理人员薪酬", path, style=2,conditions=("本期数","上年同期数"))


# 关联方应收应付款项
def addBalanceOfRelatedParties(document, path):
    addTitle(document, "(三）关联方应收应付款项", 2, True)
    addTitle(document, "1、应收关联方款项", 2, True)
    excelTableToWord(document, "应收关联方款项", path, style=2,conditions=("期末账面余额","期初账面余额"))
    addTitle(document, "2、应付关联方款项", 2, True)
    excelTableToWord(document, "应付关联方款项", path, style=2)


# 关联方承诺
def addCommitmentOfRelatedParties(document, path, context):
    companyType = context["report_params"]["companyType"]

    if companyType == "上市公司":
        df = filterDateFrame("关联方承诺上市公司",path)
        if len(df)>0:
            addTitle(document, "(四）关联方承诺", 2, True)
            dfToWord(document,df,style=2)


def addFormNote(document, path, isAll, assetsRecordsCombine, liabilitiesRecordsCombine, profitRecordsCombine,
                cashRecordsCombine,
                assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle, cashRecordsSingle,
                context):
    companyType = context["report_params"]["companyType"]
    #
    # # 计算报表项目编号
    # computeNoteNum(assetsRecordsCombine, liabilitiesRecordsCombine, profitRecordsCombine, cashRecordsCombine,
    #                assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle, cashRecordsSingle, context)

    # 添加开始
    addStart(document, context)
    # 添加报表项目
    addFsNote(document, path,context,isAll, assetsRecordsCombine, liabilitiesRecordsCombine, profitRecordsCombine,
              cashRecordsCombine,
              assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle, cashRecordsSingle)
    # 添加非报表项目
    lastNum = getLastNum(context, cashRecordsCombine, profitRecordsCombine, cashRecordsSingle, profitRecordsSingle)
    if lastNum=="":
        lastNum=0
    # 添加股份支付

    if context["noteAppend"]["shareBasedPayment"] and companyType == "国有企业":
        lastNum = lastNum + 1
        addTitle(document, "（{}）股份支付".format(to_chinese(lastNum)), 2, True)
        addShareBasedPayment(document, path, context)
    # 添加分部信息
    if context["noteAppend"]["SegmentInformation"] and companyType == "国有企业":
        lastNum = lastNum + 1
        addTitle(document, "（{}）分部信息".format(to_chinese(lastNum)), 2, True)
        addSegmentInformation(document, path)
    #     现金流量表补充资料
    lastNum = lastNum + 1
    addCashFlowReplenishment(document, lastNum, path, context)
    #     外币货币性项目
    if context["noteAppend"]["foreignCurrencyMonetaryItems"]:
        lastNum = lastNum + 1
        addForeignCurrencyMonetaryItems(document, lastNum, path)
    #     所有权或使用权受到限制的资产
    if context["noteAppend"]["limitAsset"]:
        lastNum = lastNum + 1
        addAssetsWithLimitedOwnershipOrRightOfUse(document, lastNum, path)


# 添加报表附注
# isAll显示所有的报表项目,主要用于测试环境，生产环境设置为False
def addNoteAppended(document, currentpath,parentpath,context, comparativeTable,isAll=True):
    # 公司类型
    companyType = context["report_params"]["companyType"]
    # 获取报表
    assetsRecordsCombine = searchModel(companyType, "合并", "资产表", comparativeTable)
    liabilitiesRecordsCombine = searchModel(companyType, "合并", "负债表", comparativeTable)
    profitRecordsCombine = searchModel(companyType, "合并", "利润表", comparativeTable)
    cashRecordsCombine = searchModel(companyType, "合并", "现金流量表", comparativeTable)
    assetsRecordsSingle = searchModel(companyType, "单体", "资产表", comparativeTable)
    liabilitiesRecordsSingle = searchModel(companyType, "单体", "负债表", comparativeTable)
    profitRecordsSingle = searchModel(companyType, "单体", "利润表", comparativeTable)
    cashRecordsSingle = searchModel(companyType, "单体", "现金流量表", comparativeTable)
    # 添加报表注释
    addFormNote(document, currentpath, isAll, assetsRecordsCombine, liabilitiesRecordsCombine, profitRecordsCombine,
                cashRecordsCombine,
                assetsRecordsSingle, liabilitiesRecordsSingle, profitRecordsSingle, cashRecordsSingle, context)

    # 添加其他注释
    addOtherNoteAppended(document, currentpath,parentpath, context, assetsRecordsCombine,comparativeTable)

def test():
    from project.data import testcontext
    from project.constants import comparativeTable
    from project.computeNo import computeNo
    from docx import Document
    from project.settings import setStyle

    # # 根据pandas读取的excle表格数据导入word
    MODELPATH = "D:/auditReport/project/nationalmodel.xlsx"
    # 计算附注编码
    computeNo(testcontext, comparativeTable)

    document = Document()
    # 设置中文标题
    setStyle(document)
    df = pd.DataFrame()
    addLongTermEquityInvestmentOther(document, df)
    # addInvestmentRealEstate(document, 1, MODELPATH, testcontext)
    # addNoteAppended(document, MODELPATH, MODELPATH,testcontext,comparativeTable,isAll=True)
    document.save("noteappended.docx")

if __name__ == '__main__':
    test()
